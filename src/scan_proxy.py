from __future__ import annotations

import csv
import ctypes
import io
import ipaddress
import json
import logging
import os
import re
import shutil
import signal
import socket
import sqlite3
import subprocess
import sys
import threading
import time
import uuid
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal, ROUND_HALF_UP
from html import escape
from pathlib import Path

from openpyxl import load_workbook

try:
    from pypdf import PdfReader
except ImportError:
    PdfReader = None


VERSION = "2.14.0"
# Changes:
# 2026-09-18 | Trachti | Added restricted Event 5140 fallback after Event 5145 misses.
# 2026-09-18 | Trachti | Optimized SMB open-file snapshot and pending batch resolution.
# 2026-09-18 | Trachti | Reduced CPU and I/O with cached rules, SQLite and dirty reports.
# 2026-10-22 | Trachti | Integrated English code and privacy-first public settings.




# Runtime settings are supplied by environment.local.ps1 (never committed).
BASE = (
    Path(sys.executable).resolve().parent
    if getattr(sys, "frozen", False)
    else Path(__file__).resolve().parent.parent
)
ROOT = Path(os.environ.get("SCAN_PROXY_ROOT", str(BASE / "incoming"))).expanduser()
CONFIG_PATH = Path(
    os.environ.get("SCAN_PROXY_CONFIG", str(BASE / "config" / "scan-proxy-config.xlsx"))
).expanduser()
SMB_HELPER_PATH = Path(
    os.environ.get("SCAN_PROXY_SMB_HELPER", str(BASE / "scripts" / "scan-proxy-smb.ps1"))
).expanduser()

SMB_USER = os.environ.get("SCAN_PROXY_SMB_USER", "").strip()
HOSTNAME_LOOKUPS_ENABLED = os.environ.get(
    "SCAN_PROXY_ENABLE_HOSTNAME_LOOKUPS", "false"
).strip().lower() in {"1", "true", "yes", "on"}
LOCAL_SCAN_SHARE = os.environ.get("SCAN_PROXY_LOCAL_SHARE", "").strip().lower()



FILE_SCAN_INTERVAL_SECONDS = int(os.environ.get("SCAN_PROXY_FILE_SCAN_INTERVAL_SECONDS", "5"))

REPORT_REFRESH_INTERVAL_SECONDS = int(os.environ.get("SCAN_PROXY_REPORT_REFRESH_INTERVAL_SECONDS", "600"))
PENDING_SWEEP_INTERVAL_SECONDS = int(os.environ.get("SCAN_PROXY_PENDING_SWEEP_INTERVAL_SECONDS", "60"))
PENDING_BATCH_SIZE = int(os.environ.get("SCAN_PROXY_PENDING_BATCH_SIZE", "8"))
PENDING_RETRY_DELAYS_SECONDS = (30, 120, 600, 3600, 21600, 43200)

SMB_EVENT_LOOKBACK_MINUTES = int(os.environ.get("SCAN_PROXY_SMB_EVENT_LOOKBACK_MINUTES", "10"))

SMB_SHARE_FALLBACK_MAX_AGE_SECONDS = int(os.environ.get("SCAN_PROXY_SMB_SHARE_FALLBACK_MAX_AGE_SECONDS", "30"))

SMB_IGNORED_NETWORKS = tuple(
    value.strip()
    for value in os.environ.get("SCAN_PROXY_IGNORED_NETWORKS", "").replace(",", ";").split(";")
    if value.strip()
)



LOG_ROOT = BASE / "logs"
SCANPROXY_LOG_DIR = LOG_ROOT / "scan-proxy"
ANALYTICS_LOG_DIR = LOG_ROOT / "analytics"
ANALYTICS = BASE / "analytics"
STATE = ANALYTICS / "_state"
REPORT = ANALYTICS / "analytics.html"
DB = STATE / "scan-proxy.db"
DAILY = ANALYTICS / "daily"
WEEKLY = ANALYTICS / "weekly"
MONTHLY = ANALYTICS / "monthly"
YEARLY = ANALYTICS / "yearly"

MIN_AGE = 3

MAIN_LOOP_SLEEP_SECONDS = 1

SMB_BATCH_TIMEOUT = int(os.environ.get("SCAN_PROXY_SMB_BATCH_TIMEOUT_SECONDS", "20"))
SMB_SNAPSHOT_TIMEOUT = int(os.environ.get("SCAN_PROXY_SMB_SNAPSHOT_TIMEOUT_SECONDS", "4"))
SMB_SNAPSHOT_TTL_SECONDS = 8

PENDING_MAX_HOURS = int(os.environ.get("SCAN_PROXY_PENDING_MAX_HOURS", "48"))

SMB_OPEN_CACHE_SECONDS = 2
SMB_RESOLVED_CACHE_SECONDS = 300
SMB_MISS_CACHE_SECONDS = 30
SMB_EXACT_HINT_SECONDS = 3600

HOSTNAME_LOOKUP_TIMEOUT_SECONDS = 2
HOSTNAME_SUCCESS_TTL_SECONDS = 86400
HOSTNAME_FAILURE_TTL_SECONDS = 21600
HOSTNAME_BATCH_SIZE = 4

TARGET_DIR_CACHE_TTL_SECONDS = 600

PERFORMANCE_LOG_INTERVAL_SECONDS = 300

DAILY_RETENTION_DAYS = int(os.environ.get("SCAN_PROXY_DAILY_RETENTION_DAYS", "90"))

MOVE_RETRY_SECONDS = 5

PENDING_CACHE_SECONDS = 5
SMB_IGNORED_NETS = tuple(
    ipaddress.ip_network(network, strict=False)
    for network in SMB_IGNORED_NETWORKS
)
STOP = threading.Event()

LOCAL_SCAN_SERVER_ALIASES = set()
if os.environ.get("COMPUTERNAME"):
    LOCAL_SCAN_SERVER_ALIASES.add(os.environ["COMPUTERNAME"].split(".", 1)[0].lower())
LOCAL_SCAN_SERVER_ALIASES.add(socket.gethostname().split(".", 1)[0].lower())
LOCAL_SCAN_SERVER_ALIASES.update(
    alias.strip().split(".", 1)[0].lower()
    for alias in os.environ.get("SCAN_PROXY_LOCAL_SERVER_ALIASES", "").replace(",", ";").split(";")
    if alias.strip()
)

_TARGET_DIR_CACHE = {}

@dataclass(frozen=True, slots=True)
class Rule:
    row: int
    code: str
    source_excel: str
    source: Path
    target: Path
    mode: str
    action: str
    scanner: str

@dataclass(slots=True)
class RuleGroup:
    source: Path
    source_excel: str
    identifiers: dict = field(default_factory=dict)
    subfolders: list = field(default_factory=list)
    error_until: float = 0.0

@dataclass(frozen=True, slots=True)
class FileCandidate:
    path: Path
    stat: object


def resolve_path(value):
    path = Path(str(value or ".").strip())
    if path.is_absolute():
        return path
    return ROOT / path


def localize_scan_source(path):
    """Map only the explicitly configured local SMB share to ROOT; Excel paths stay unchanged."""
    if os.name != "nt":
        return path
    text = str(path).strip().replace("/", "\\")
    match = re.match(r"^\\\\([^\\]+)\\([^\\]+)(?:\\(.*))?$", text)
    if not match:
        return path
    server, share, rest = match.groups()
    server = server.split(".", 1)[0].lower()
    if not LOCAL_SCAN_SHARE or server not in LOCAL_SCAN_SERVER_ALIASES or share.lower() != LOCAL_SCAN_SHARE:
        return path
    if not rest:
        return ROOT
    return ROOT.joinpath(*[part for part in rest.split("\\") if part])


def norm_path_key(path):
    return os.path.normcase(os.path.normpath(str(path)))


def norm_relative(value):
    return str(value or "").strip().replace("/", "\\").strip("\\").lower()


def ensure_target_dir(path):
    """Avoid repeated mkdir calls for the same destination directory."""
    now = time.monotonic()
    key = norm_path_key(path)
    if _TARGET_DIR_CACHE.get(key, 0.0) > now:
        return
    path.mkdir(parents=True, exist_ok=True)
    _TARGET_DIR_CACHE[key] = now + TARGET_DIR_CACHE_TTL_SECONDS
    if len(_TARGET_DIR_CACHE) > 2048:
        _TARGET_DIR_CACHE.clear()

def is_enabled(value):
    return str(value).strip().lower() in {
        "true",
        "on",
        "1",
        "enabled",
        "yes",
        "x",
    }

def normalize_ip(value):
    text = str(value or "").strip().strip("[]")
    try:
        addr = ipaddress.ip_address(text)
        if isinstance(addr, ipaddress.IPv6Address) and addr.ipv4_mapped:
            addr = addr.ipv4_mapped
        if str(addr) in {
            "127.0.0.1",
            "::1",
            "0.0.0.0",
            "::",
        }:
            return ""

        if any(
            addr.version == network.version and addr in network
            for network in SMB_IGNORED_NETS
        ):
            return ""
        return str(addr)
    except ValueError:
        return ""

def pending_retry_delay(attempts):
    """Return the backoff interval for a scanner with unresolved client IP."""
    index = min(max(int(attempts), 0), len(PENDING_RETRY_DELAYS_SECONDS) - 1)
    return PENDING_RETRY_DELAYS_SECONDS[index]


def reverse_lookup(ip):
    """Resolve an IP via nslookup, with a socket-based fallback."""
    client = normalize_ip(ip)
    if not client:
        return "", "None"

    executable = "nslookup.exe" if os.name == "nt" else shutil.which("nslookup")
    if executable:
        try:
            result = subprocess.run(
                [executable, client],
                capture_output=True,
                text=True,
                encoding="utf-8",
                errors="replace",
                timeout=HOSTNAME_LOOKUP_TIMEOUT_SECONDS,
                creationflags=(
                    getattr(subprocess, "CREATE_NO_WINDOW", 0)
                    | getattr(subprocess, "BELOW_NORMAL_PRIORITY_CLASS", 0)
                ),
            )
            text = "\n".join((result.stdout or "", result.stderr or ""))
            matches = re.findall(
                r"(?im)^\s*(?:name|name)\s*[:=]\s*(\S.*?)\s*$",
                text,
            )
            for value in matches:
                hostname = value.strip().rstrip(".")
                if hostname and not normalize_ip(hostname) and hostname != client:
                    return hostname, "nslookup"
        except (OSError, subprocess.SubprocessError):
            pass
        return "", "None"

    try:
        hostname = socket.gethostbyaddr(client)[0].strip().rstrip(".")
        if hostname and hostname != client:
            return hostname, "socket"
    except (OSError, socket.herror, socket.gaierror):
        pass
    return "", "None"


def utc_now():
    return datetime.now(timezone.utc).isoformat().replace("+00:00", "Z")

def parse_utc(value):
    try:
        return datetime.fromisoformat(
            str(value).replace("Z", "+00:00")
        ).astimezone(timezone.utc)
    except (ValueError, TypeError):
        return None


def mb(size):
    value = Decimal(int(size or 0)) / Decimal(1048576)
    rounded = value.quantize(
        Decimal("0.01"),
        rounding=ROUND_HALF_UP,
    )
    return f"{rounded:.2f}"


def pages(file, log):
    if file.suffix.lower() != ".pdf":
        return 1
    try:
        if PdfReader:
            reader = PdfReader(
                str(file),
                strict=False,
            )
            return max(1, len(reader.pages))
        content = file.read_bytes()
        return max(
            1,
            len(re.findall(rb"/Type\s*/Page\b", content)),
        )
    except Exception:
        log.warning(
            "ANALYTICS WARNING | "
            "Could not determine PDF page count",
        )
        return 1


class DailyFileLogHandler(logging.Handler):
    def __init__(self, folder):
        super().__init__()
        self.folder = folder
        self.day = None
        self.stream = None
        folder.mkdir(
            parents=True,
            exist_ok=True,
        )

    def emit(self, record):
        try:
            day = datetime.now().strftime("%Y-%m-%d")
            if day != self.day:
                if self.stream:
                    self.stream.close()
                log_file = self.folder / f"{day}-scan-proxy.log"
                self.stream = log_file.open(
                    "a",
                    encoding="utf-8",
                )
                self.day = day
            self.stream.write(
                self.format(record) + "\n"
            )
            self.stream.flush()
        except Exception:
            self.handleError(record)
    def close(self):
        if self.stream:
            self.stream.close()
        super().close()

def get_logger(name, folder):
    log = logging.getLogger(name)
    log.setLevel(logging.INFO)
    log.propagate = False
    if not log.handlers:
        handler = DailyFileLogHandler(folder)
        handler.setFormatter(
            logging.Formatter(
                "%(asctime)s | %(levelname)s | %(message)s"
            )
        )
        log.addHandler(handler)
    return log


def get_logs():
    scanproxy_log = get_logger(
        "scanproxy",
        SCANPROXY_LOG_DIR,
    )
    analytics_log = get_logger(
        "analytics",
        ANALYTICS_LOG_DIR,
    )
    return scanproxy_log, analytics_log




def load_rules():
    workbook = load_workbook(
        CONFIG_PATH,
        read_only=True,
        data_only=True,
    )
    try:
        worksheets = sorted(
            workbook.worksheets,
            key=lambda sheet: sheet.title.strip().lower() != "rules",
        )
        for worksheet in worksheets:
            header = next(
                worksheet.iter_rows(
                    min_row=1,
                    max_row=1,
                    values_only=True,
                ),
                (),
            )
            columns = {
                str(value).strip().lower(): index
                for index, value in enumerate(header)
                if value is not None
            }
            required = {
                "identifier",
                "source",
                "destination",
                "mode",
            }
            has_active_column = (
                "enabled" in columns
                or "status" in columns
            )
            if required <= columns.keys() and has_active_column:
                break
        else:
            raise ValueError(
                "No worksheet contains Identifier, Source, Destination, Mode "
                "and Enabled/Status columns"
            )
        columns["enabled"] = columns.get(
            "enabled",
            columns.get("status"),
        )
        def val(row, name, default=None):
            index = columns.get(name)
            if index is not None and index < len(row):
                return row[index]
            return default
        rules = []
        for row_number, row in enumerate(
            worksheet.iter_rows(
                min_row=2,
                values_only=True,
            ),
            2,
        ):
            if not any(value not in (None, "") for value in row):
                continue
            rules.append(
                (
                    row_number,
                    val(row, "identifier"),
                    val(row, "source"),
                    val(row, "destination"),
                    val(row, "mode"),
                    val(row, "enabled"),
                    val(row, "action", "move") or "move",
                    val(row, "scanner"),
                )
            )
        return worksheet.title, rules
    finally:
        workbook.close()



def compile_rules(raw_rules, log):
    """Validate and group Excel rules once when the workbook is loaded."""
    groups = {}
    active_count = 0
    mapped_sources = set()
    seen_rules = {}

    for row, code, source, target, mode, enabled, action, scanner in raw_rules:
        if not is_enabled(enabled):
            continue
        if not source or not target:
            log.error("RULE ERROR | Row=%s | Source/Destination missing", row)
            continue

        source_excel = str(source).strip()
        src_excel_path = resolve_path(source_excel)
        src = localize_scan_source(src_excel_path)
        dst = resolve_path(target)
        # The administrator controls the workbook; validate malformed paths early.
        if str(src).strip() == str(dst).strip():
            log.error("RULE ERROR | Source equals destination | Row=%s", row)
            continue
        mode = str(mode or "").strip().lower()
        action = str(action or "move").strip().lower()
        code_text = str(code or "").strip()

        if action not in {"move", "copy"}:
            log.error("RULE ERROR | Row=%s | Invalid action=%s", row, action)
            continue
        if mode not in {"identifier", "subfolder", "subfolders"}:
            log.error("RULE ERROR | Row=%s | Mode=%s", row, mode)
            continue
        if mode == "identifier" and not code_text:
            log.error("RULE ERROR | Row=%s | Identifier missing", row)
            continue

        if mode in {"subfolder", "subfolders"}:
            dst_io = localize_scan_source(dst)
            try:
                if dst_io == src or dst_io.is_relative_to(src):
                    log.error("RULE ERROR | Destination is inside source | Row=%s", row)
                    continue
            except (OSError, ValueError):
                pass

        signature = (
            code_text.lower(), norm_path_key(src), norm_path_key(dst),
            mode, action,
        )
        if signature in seen_rules:
            log.warning(
                "DUPLICATE RULE SKIPPED | Row=%s | DuplicateOf=%s",
                row, seen_rules[signature],
            )
            continue
        seen_rules[signature] = row

        rule = Rule(
            row=row,
            code=code_text,
            source_excel=source_excel,
            source=src,
            target=dst,
            mode=mode,
            action=action,
            scanner=str(scanner or "").strip(),
        )
        key = norm_path_key(src)
        group = groups.setdefault(key, RuleGroup(src, source_excel))
        if mode == "identifier":
            group.identifiers.setdefault(code_text.lower(), []).append(rule)
        else:
            group.subfolders.append(rule)
        active_count += 1

        if norm_path_key(src) != norm_path_key(src_excel_path) and key not in mapped_sources:
            mapped_sources.add(key)
            log.info("LOCAL SOURCE OPTIMIZED")

    return list(groups.values()), active_count


def scan_source(root, recursive):
    """Walk a source directory with os.scandir and cache file metadata."""
    files = []
    stack = [root]
    while stack:
        folder = stack.pop()
        with os.scandir(folder) as entries:
            for entry in entries:
                if entry.name.lower().endswith(".part"):
                    continue
                try:
                    if entry.is_file(follow_symlinks=False):
                        files.append(FileCandidate(Path(entry.path), entry.stat(follow_symlinks=False)))
                    elif recursive and entry.is_dir(follow_symlinks=False):
                        stack.append(Path(entry.path))
                except OSError:
                    continue
    return files




class SmbResolver:
    """Resolve SMB clients using a low-cost open-file snapshot and a historical batch."""

    def __init__(self, log):
        self.log = log
        self.exact_clients = {}
        self.snapshot_paths = {}
        self.snapshot_relative = {}
        self.snapshot_until = 0.0
        self.snapshot_wall_time = 0.0

    @staticmethod
    def empty(reason="No unambiguous SMB client IP attribution"):
        return {
            "client": "",
            "resolved": False,
            "source": "None",
            "reason": reason,
            "record_id": 0,
            "is_open": False,
            "confidence": "None",
        }

    def refresh_open_snapshot(self, force=False):
        """Run at most one Get-SmbOpenFile process for multiple files."""
        started = time.perf_counter()
        now = time.monotonic()
        if not force and self.snapshot_until > now:
            return 0.0
        self.snapshot_paths = {}
        self.snapshot_relative = {}
        self.snapshot_until = now + SMB_SNAPSHOT_TTL_SECONDS
        self.snapshot_wall_time = time.time()
        if not SMB_USER or os.name != "nt" or not SMB_HELPER_PATH.is_file():
            return 0.0

        command = [
            "powershell.exe", "-NoLogo", "-NoProfile", "-NonInteractive",
            "-ExecutionPolicy", "Bypass", "-File", str(SMB_HELPER_PATH),
            "-Mode", "OpenSnapshot", "-UserName", SMB_USER,
            "-IgnoredNetworks", ";".join(SMB_IGNORED_NETWORKS),
        ]
        try:
            proc = subprocess.run(
                command,
                capture_output=True,
                text=True,
                encoding="utf-8",
                errors="replace",
                timeout=SMB_SNAPSHOT_TIMEOUT,
                creationflags=(
                    getattr(subprocess, "CREATE_NO_WINDOW", 0)
                    | getattr(subprocess, "BELOW_NORMAL_PRIORITY_CLASS", 0)
                ),
            )
            if proc.returncode:
                raise RuntimeError((proc.stderr or f"ExitCode={proc.returncode}").strip())
            lines = [line.lstrip("\ufeff").strip() for line in proc.stdout.splitlines() if line.strip()]
            data = json.loads(lines[-1]) if lines else []
            if isinstance(data, dict):
                data = [data]
            for item in data:
                client = normalize_ip(item.get("Client"))
                if not client:
                    continue
                file_path = item.get("Path")
                relative = norm_relative(item.get("RelativePath"))
                if file_path:
                    self.snapshot_paths.setdefault(norm_path_key(file_path), set()).add(client)
                if relative:
                    self.snapshot_relative.setdefault(relative, set()).add(client)
        except Exception as exc:
            self.log.warning("SMB SNAPSHOT ERROR | ErrorType=%s", type(exc).__name__)
        return (time.perf_counter() - started) * 1000.0

    def resolve(self, file, relative="", mtime="", reference=""):
        """Match only the in-memory SMB snapshot; never launch one process per file."""
        path_key = norm_path_key(file)
        rel_key = norm_relative(relative)
        now = time.monotonic()

        clients = self.snapshot_paths.get(path_key)
        if not clients and rel_key:
            clients = self.snapshot_relative.get(rel_key)
        if clients:
            if len(clients) == 1:
                client = next(iter(clients))
                result = {
                    "client": client,
                    "resolved": True,
                    "source": "Get-SmbOpenFile-Snapshot",
                    "reason": "Exact SMB open file found in snapshot",
                    "record_id": 0,
                    "is_open": True,
                    "confidence": "ExactOpenFile",
                }
                cached = dict(result)
                cached["is_open"] = False
                cached["source"] = "Get-SmbOpenFile-Cache"
                cached["reason"] = "Previously observed exact SMB open file"
                cached["confidence"] = "ExactOpenFileCached"
                self.exact_clients[(path_key, mtime)] = (cached, now + SMB_EXACT_HINT_SECONDS)
                return result
            result = self.empty("Ambiguous SMB open-file snapshot")
            result["is_open"] = True
            result["confidence"] = "Ambiguous"
            return result

        hint = self.exact_clients.get((path_key, mtime))
        if hint:
            if hint[1] > now:
                return dict(hint[0])
            self.exact_clients.pop((path_key, mtime), None)
        return self.empty()

    def resolve_batch(self, rows):
        """Resolve historical pending files with one PowerShell process per batch."""
        rows = list(rows)
        if not rows:
            return {}
        empty = self.empty()
        if not SMB_USER or os.name != "nt" or not SMB_HELPER_PATH.is_file():
            return {str(row["id"]): dict(empty) for row in rows}

        payload = [
            {
                "Id": str(row["id"]),
                "FilePath": str(row["original_path"]),
                "RelativePath": str(row["relative_path"] or ""),
                "ReferenceTimeUtc": str(row["file_mtime_utc"] or row["completed_at_utc"] or utc_now()),
            }
            for row in rows
        ]
        command = [
            "powershell.exe", "-NoLogo", "-NoProfile", "-NonInteractive",
            "-ExecutionPolicy", "Bypass", "-File", str(SMB_HELPER_PATH),
            "-Mode", "ResolveBatch", "-UserName", SMB_USER,
            "-LookbackMinutes", str(SMB_EVENT_LOOKBACK_MINUTES),
            "-SessionFallbackMaxAgeSeconds", str(SMB_SHARE_FALLBACK_MAX_AGE_SECONDS),
            "-IgnoredNetworks", ";".join(SMB_IGNORED_NETWORKS),
        ]
        try:
            proc = subprocess.run(
                command,
                input=json.dumps(payload, ensure_ascii=False),
                capture_output=True,
                text=True,
                encoding="utf-8",
                errors="replace",
                timeout=SMB_BATCH_TIMEOUT,
                creationflags=(
                    getattr(subprocess, "CREATE_NO_WINDOW", 0)
                    | getattr(subprocess, "BELOW_NORMAL_PRIORITY_CLASS", 0)
                ),
            )
            if proc.returncode:
                raise RuntimeError((proc.stderr or f"ExitCode={proc.returncode}").strip())
            output = [line.lstrip("\ufeff").strip() for line in proc.stdout.splitlines() if line.strip()]
            data = json.loads(output[-1]) if output else []
            if isinstance(data, dict):
                data = [data]
            results = {}
            for item in data:
                row_id = str(item.get("Id") or "").strip()
                if not row_id:
                    continue
                client = normalize_ip(item.get("Client"))
                results[row_id] = {
                    "client": client,
                    "resolved": bool(client),
                    "source": str(item.get("Source") or "None"),
                    "reason": str(item.get("Reason") or empty["reason"]),
                    "record_id": int(item.get("RecordId") or 0),
                    "is_open": bool(item.get("Open")),
                    "confidence": str(item.get("Confidence") or "None"),
                }
            for row in rows:
                results.setdefault(str(row["id"]), dict(empty))
            return results
        except Exception as exc:
            self.log.warning("SMB BATCH ERROR | Files=%s | ErrorType=%s", len(rows), type(exc).__name__)
            failed = self.empty(f"SMB batch failed: {exc}")
            return {str(row["id"]): dict(failed) for row in rows}

    def forget(self, file):
        path_key = norm_path_key(file)
        for cache_key in list(self.exact_clients):
            if cache_key[0] == path_key:
                self.exact_clients.pop(cache_key, None)




class Analytics:


    def __init__(self, log):
        self.log = log
        self.last_move_retry = 0.0
        self._pending_cache = set()
        self._pending_cache_until = 0.0
        self.db = None
        self._reports_dirty = True
        self._dirty_dates = set()
        self._last_maintenance_day = None

    def connect(self):
        """Reuse one SQLite connection for this single-threaded service."""
        if self.db is None:
            self.db = sqlite3.connect(DB, timeout=10)
            self.db.row_factory = sqlite3.Row
            self.db.execute("PRAGMA busy_timeout=10000")
            self.db.execute("PRAGMA synchronous=NORMAL")
            self.db.execute("PRAGMA temp_store=MEMORY")
        return self.db

    def close(self):
        if self.db is not None:
            try:
                self.db.close()
            finally:
                self.db = None

    def mark_dirty(self, completed_date=None):
        self._reports_dirty = True
        if completed_date:
            self._dirty_dates.add(str(completed_date))

    def init(self):
        for folder in (
            ANALYTICS,
            STATE,
            DAILY,
            WEEKLY,
            MONTHLY,
            YEARLY,
        ):
            folder.mkdir(
                parents=True,
                exist_ok=True,
            )
        with self.connect() as db:
            db.execute(
                "PRAGMA journal_mode=WAL"
            )
            db.executescript(
                """
                CREATE TABLE IF NOT EXISTS scans (
                    id TEXT PRIMARY KEY,
                    completed_at_utc TEXT NOT NULL,
                    completed_at_local TEXT NOT NULL,
                    completed_date TEXT NOT NULL,
                    original_path TEXT NOT NULL,
                    relative_path TEXT NOT NULL,
                    completed_path TEXT NOT NULL,
                    routing_label TEXT NOT NULL,
                    file_mtime_utc TEXT NOT NULL,
                    pages INTEGER NOT NULL,
                    size_bytes INTEGER NOT NULL,
                    scanner_ip TEXT,
                    resolution_source TEXT NOT NULL,
                    resolution_reason TEXT NOT NULL,
                    resolution_confidence TEXT NOT NULL,
                    resolution_record_id INTEGER NOT NULL DEFAULT 0,
                    resolution_state TEXT NOT NULL,
                    resolve_attempts INTEGER NOT NULL DEFAULT 0,
                    next_resolve_epoch REAL NOT NULL DEFAULT 0,
                    resolve_deadline_utc TEXT NOT NULL,
                    created_at_utc TEXT NOT NULL,
                    updated_at_utc TEXT NOT NULL
                );

                CREATE INDEX IF NOT EXISTS idx_scans_pending
                    ON scans (
                        resolution_state,
                        next_resolve_epoch
                    );

                CREATE INDEX IF NOT EXISTS idx_scans_report
                    ON scans (
                        completed_date,
                        scanner_ip,
                        resolution_state
                    );

                CREATE INDEX IF NOT EXISTS idx_scans_resolved_report
                    ON scans (
                        resolution_state,
                        completed_date,
                        scanner_ip
                    );

                CREATE TABLE IF NOT EXISTS dns_cache (
                    scanner_ip TEXT PRIMARY KEY,
                    hostname TEXT NOT NULL DEFAULT '',
                    source TEXT NOT NULL DEFAULT 'None',
                    last_lookup_utc TEXT NOT NULL,
                    next_lookup_epoch REAL NOT NULL DEFAULT 0,
                    failures INTEGER NOT NULL DEFAULT 0
                );

                CREATE INDEX IF NOT EXISTS idx_dns_cache_due
                    ON dns_cache (next_lookup_epoch);

                CREATE TABLE IF NOT EXISTS pending_moves (
                    id TEXT PRIMARY KEY,
                    source_norm TEXT NOT NULL,
                    source_path TEXT NOT NULL,
                    target_path TEXT NOT NULL,
                    relative_path TEXT NOT NULL,
                    routing_label TEXT NOT NULL,
                    source_size INTEGER NOT NULL,
                    source_mtime_ns INTEGER NOT NULL,
                    file_mtime_utc TEXT NOT NULL,
                    committed_at_utc TEXT NOT NULL,
                    resolution_json TEXT NOT NULL,
                    reason TEXT NOT NULL,
                    state TEXT NOT NULL,
                    created_at_utc TEXT NOT NULL,
                    last_attempt_utc TEXT NOT NULL
                );

                CREATE INDEX IF NOT EXISTS idx_pending_moves_source
                    ON pending_moves (
                        source_norm,
                        state
                    );
                """
            )
        self.refresh_reports(force=True)


    def pending_move(self, source):
        now = time.monotonic()
        if now >= self._pending_cache_until:
            with self.connect() as db:
                rows = db.execute(
                    """
                    SELECT source_norm
                    FROM pending_moves
                    WHERE state IN ('pending', 'blocked')
                    """
                ).fetchall()
            self._pending_cache = {
                row["source_norm"]
                for row in rows
            }
            self._pending_cache_until = (
                now + PENDING_CACHE_SECONDS
            )
        key = os.path.normcase(
            os.path.normpath(
                str(source)
            )
        )
        return key in self._pending_cache
    def add_move(
        self,
        source,
        target,
        relative,
        label,
        stat,
        mtime,
        resolution,
        reason,
    ):
        

        now = utc_now()
        move_id = uuid.uuid4().hex
        source_norm = os.path.normcase(
            os.path.normpath(
                str(source)
            )
        )
        values = (
            move_id,
            source_norm,
            str(source),
            str(target),
            str(relative),
            label,
            int(stat.st_size),
            int(stat.st_mtime_ns),
            mtime,
            now,
            json.dumps(resolution),
            type(reason).__name__,
            "pending",
            now,
            now,
        )
        with self.connect() as db:
            db.execute(
                """
                INSERT INTO pending_moves
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                values,
            )
        self._pending_cache.add(source_norm)
    def add_scan(
        self,
        source,
        relative,
        target,
        label,
        mtime,
        resolution,
        scan_id=None,
        completed=None,
        page_count=None,
        size_bytes=None,
    ):
        

        done = (
            parse_utc(completed)
            or datetime.now(timezone.utc)
        )
        local = done.astimezone()
        now = utc_now()

        client = normalize_ip(
            (resolution or {}).get("client")
        )
        # Avoid retaining document paths when SMB attribution is disabled.
        state = "resolved" if client else ("pending" if SMB_USER else "unresolved")
        deadline = (
            datetime.now(timezone.utc)
            + timedelta(hours=PENDING_MAX_HOURS)
        ).isoformat().replace(
            "+00:00",
            "Z",
        )
        values = (
            scan_id or uuid.uuid4().hex,
            done.isoformat().replace("+00:00", "Z"),
            local.isoformat(timespec="seconds"),
            local.date().isoformat(),
            str(source) if state == "pending" else "",
            str(relative) if state == "pending" else "",
            "",
            "",
            mtime,
            int(page_count) if page_count is not None else pages(target, self.log),
            int(size_bytes) if size_bytes is not None else target.stat().st_size,
            client or None,
            str(
                (resolution or {}).get("source")
                or "None"
            ),
            str(
                (resolution or {}).get("reason")
                or ""
            ),
            str(
                (resolution or {}).get("confidence")
                or "None"
            ),
            int(
                (resolution or {}).get("record_id")
                or 0
            ),
            state,
            0,
            time.time() + pending_retry_delay(0),
            deadline,
            now,
            now,
        )
        with self.connect() as db:
            db.execute(
                """
                INSERT OR IGNORE INTO scans
                VALUES (
                    ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?,
                    ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?
                )
                """,
                values,
            )
        self.mark_dirty(local.date().isoformat())
        self.log.info("ANALYTICS | State=%s | Pages=%s", state, values[9])


    def retry_moves(self):
        now = time.monotonic()
        if now - self.last_move_retry < MOVE_RETRY_SECONDS:
            return
        self.last_move_retry = now
        

        self._pending_cache_until = 0.0
        with self.connect() as db:
            rows = db.execute(
                """
                SELECT *
                FROM pending_moves
                WHERE state = 'pending'
                ORDER BY created_at_utc
                LIMIT 10
                """
            ).fetchall()
        for row in rows:
            src = Path(row["source_path"])
            dst = Path(row["target_path"])
            if not dst.exists():
                with self.connect() as db:
                    db.execute(
                        "DELETE FROM pending_moves WHERE id = ?",
                        (row["id"],),
                    )
                continue
            try:
                stat = src.stat()
            except FileNotFoundError:
                stat = None
            except OSError:
                continue
            source_changed = (
                stat
                and (
                    stat.st_size != row["source_size"]
                    or stat.st_mtime_ns != row["source_mtime_ns"]
                )
            )
            if source_changed:
                try:
                    dst.unlink()
                    state = None
                except OSError:
                    state = "blocked"
                with self.connect() as db:
                    if state:
                        db.execute(
                            """
                            UPDATE pending_moves
                            SET state = 'blocked',
                                last_attempt_utc = ?
                            WHERE id = ?
                            """,
                            (
                                utc_now(),
                                row["id"],
                            ),
                        )
                    else:
                        db.execute(
                            "DELETE FROM pending_moves WHERE id = ?",
                            (row["id"],),
                        )
                continue
            if stat:
                try:
                    src.unlink()
                except OSError:
                    continue
            try:
                resolution = json.loads(
                    row["resolution_json"] or "{}"
                )
            except json.JSONDecodeError:
                resolution = {}
            self.add_scan(
                src,
                row["relative_path"],
                dst,
                row["routing_label"],
                row["file_mtime_utc"],
                resolution,
                row["id"],
                row["committed_at_utc"],
            )
            with self.connect() as db:
                db.execute(
                    "DELETE FROM pending_moves WHERE id = ?",
                    (row["id"],),
                )


    def resolve_pending(self, smb):
        """Resolve due pending IPs in a single low-CPU batch."""
        now = time.time()
        now_utc = datetime.now(timezone.utc)
        with self.connect() as db:
            rows = db.execute(
                """
                SELECT * FROM scans
                WHERE resolution_state = 'pending' AND next_resolve_epoch <= ?
                ORDER BY next_resolve_epoch, completed_at_utc
                LIMIT ?
                """,
                (now, PENDING_BATCH_SIZE),
            ).fetchall()
        if not rows:
            return 0

        expired, active_rows = [], []
        for row in rows:
            deadline = parse_utc(row["resolve_deadline_utc"])
            (expired if deadline and deadline <= now_utc else active_rows).append(row)

        changed = 0
        if expired:
            stamp = utc_now()
            with self.connect() as db:
                db.executemany(
                    """UPDATE scans SET resolution_state='unresolved',
                        original_path='', relative_path='', completed_path='', routing_label='',
                        updated_at_utc=? WHERE id=?""",
                    [(stamp, row["id"]) for row in expired],
                )
            for row in expired:
                self.mark_dirty(row["completed_date"])
                changed += 1
                self.log.warning("SMB ATTRIBUTION EXPIRED | Attempts=%s", row["resolve_attempts"])

        if not active_rows:
            return changed

        results = smb.resolve_batch(active_rows)
        resolved_logs = []
        stamp = utc_now()
        retry_now = time.time()
        with self.connect() as db:
            for row in active_rows:
                result = results.get(str(row["id"]), {})
                client = normalize_ip(result.get("client"))
                attempts = int(row["resolve_attempts"] or 0) + 1
                reason = str(result.get("reason") or "No unambiguous SMB client IP attribution")[:1000]
                if client:
                    db.execute(
                        """
                        UPDATE scans SET scanner_ip=?, resolution_source=?, resolution_reason=?,
                            resolution_confidence=?, resolution_record_id=?, resolution_state='resolved',
                            original_path='', relative_path='', completed_path='', routing_label='',
                            resolve_attempts=?, updated_at_utc=? WHERE id=?
                        """,
                        (client, str(result.get("source") or "None"), reason,
                         str(result.get("confidence") or "None"), int(result.get("record_id") or 0),
                         attempts, stamp, row["id"]),
                    )
                    self.mark_dirty(row["completed_date"])
                    changed += 1
                    resolved_logs.append(str(result.get("source") or "None"))
                else:
                    db.execute(
                        """
                        UPDATE scans SET resolve_attempts=?, next_resolve_epoch=?,
                            resolution_reason=?, updated_at_utc=? WHERE id=?
                        """,
                        (attempts, retry_now + pending_retry_delay(attempts), reason, stamp, row["id"]),
                    )
        for source in resolved_logs:
            self.log.info("SMB RESOLVED | Evidence=%s", source)
        return changed


    def refresh_hostnames(self, ips):
        if not HOSTNAME_LOOKUPS_ENABLED:
            return set()
        candidates = sorted({normalize_ip(ip) for ip in ips if normalize_ip(ip)})
        changed_ips = set()
        if not candidates:
            return changed_ips

        with self.connect() as db:
            cached = {
                row["scanner_ip"]: row
                for row in db.execute(
                    "SELECT * FROM dns_cache"
                ).fetchall()
            }

        now = time.time()
        due = [
            ip for ip in candidates
            if ip not in cached or float(cached[ip]["next_lookup_epoch"] or 0) <= now
        ][:HOSTNAME_BATCH_SIZE]
        if not due:
            return changed_ips

        results = [reverse_lookup(ip) for ip in due]

        with self.connect() as db:
            for ip, (hostname, source) in zip(due, results):
                previous = cached.get(ip)
                previous_name = str(previous["hostname"] or "") if previous else ""
                failures = 0 if hostname else int(previous["failures"] or 0) + 1 if previous else 1
                effective_name = hostname or previous_name
                ttl = (
                    HOSTNAME_SUCCESS_TTL_SECONDS
                    if hostname
                    else HOSTNAME_FAILURE_TTL_SECONDS
                )
                db.execute(
                    """
                    INSERT INTO dns_cache (
                        scanner_ip, hostname, source, last_lookup_utc,
                        next_lookup_epoch, failures
                    ) VALUES (?, ?, ?, ?, ?, ?)
                    ON CONFLICT(scanner_ip) DO UPDATE SET
                        hostname = excluded.hostname,
                        source = excluded.source,
                        last_lookup_utc = excluded.last_lookup_utc,
                        next_lookup_epoch = excluded.next_lookup_epoch,
                        failures = excluded.failures
                    """,
                    (
                        ip,
                        effective_name,
                        source if hostname else (previous["source"] if previous else "None"),
                        utc_now(),
                        now + ttl,
                        failures,
                    ),
                )
                if hostname and hostname != previous_name:
                    changed_ips.add(ip)
                    self.log.info("HOSTNAME RESOLVED | Source=%s", source)
        return changed_ips

    def hostnames(self):
        if not HOSTNAME_LOOKUPS_ENABLED:
            return {}
        with self.connect() as db:
            rows = db.execute(
                "SELECT scanner_ip, hostname FROM dns_cache WHERE hostname <> ''"
            ).fetchall()
        return {row["scanner_ip"]: row["hostname"] for row in rows}

    def quality_data(self):
        """Return daily counts of resolved, pending and unresolved scan records."""
        with self.connect() as db:
            rows = db.execute(
                """
                SELECT
                    completed_date,
                    COUNT(*) AS total,
                    COALESCE(SUM(pages), 0) AS pages,
                    COALESCE(SUM(size_bytes), 0) AS bytes,
                    MIN(completed_at_local) AS first_scan,
                    MAX(completed_at_local) AS last_scan,
                    SUM(CASE WHEN resolution_state = 'resolved' THEN 1 ELSE 0 END) AS resolved,
                    SUM(CASE WHEN resolution_state = 'pending' THEN 1 ELSE 0 END) AS pending,
                    SUM(CASE WHEN resolution_state = 'unresolved' THEN 1 ELSE 0 END) AS unresolved
                FROM scans
                GROUP BY completed_date
                ORDER BY completed_date
                """
            ).fetchall()
        return [dict(row) for row in rows]

    def data(self):
        with self.connect() as db:
            rows = db.execute(
                """
                SELECT
                    completed_date,
                    scanner_ip,
                    COUNT(*) AS files,
                    COALESCE(SUM(pages), 0) AS pages,
                    COALESCE(SUM(size_bytes), 0) AS bytes,
                    MIN(completed_at_local) AS first_scan,
                    MAX(completed_at_local) AS last_scan
                FROM scans
                WHERE resolution_state = 'resolved'
                  AND scanner_ip IS NOT NULL
                GROUP BY
                    completed_date,
                    scanner_ip
                ORDER BY
                    completed_date,
                    scanner_ip
                """
            ).fetchall()
        return [
            row
            for row in rows
            if normalize_ip(row["scanner_ip"])
        ]
    @staticmethod
    def write_csv(path, fields, rows):
        """Write a CSV atomically only when its content has changed."""
        path.parent.mkdir(parents=True, exist_ok=True)
        buffer = io.StringIO(newline="")
        writer = csv.DictWriter(
            buffer,
            fieldnames=fields,
            delimiter=";",
            extrasaction="ignore",
        )
        writer.writeheader()
        writer.writerows(rows)
        payload = buffer.getvalue().encode("utf-8-sig")

        try:
            if path.read_bytes() == payload:
                return False
        except FileNotFoundError:
            pass

        tmp = path.with_suffix(path.suffix + ".tmp")
        with tmp.open("wb") as output:
            output.write(payload)
            output.flush()
            os.fsync(output.fileno())
        os.replace(tmp, path)
        return True

    @staticmethod
    def write_text(path, text):
        """Atomically write changed text while avoiding unnecessary I/O."""
        payload = text.encode("utf-8")
        try:
            if path.read_bytes() == payload:
                return False
        except FileNotFoundError:
            pass
        tmp = path.with_suffix(path.suffix + ".tmp")
        with tmp.open("wb") as output:
            output.write(payload)
            output.flush()
            os.fsync(output.fileno())
        os.replace(tmp, path)
        return True
    def build_report_html(
        self,
        embedded,
        latest,
        interval_label,
    ):
        """Build standalone local analytics HTML with embedded data."""
        return f"""<!doctype html>
<html lang="en">
<head>
    <meta charset="utf-8">
    <meta name="viewport" content="width=device-width, initial-scale=1">
    <title>Scan Proxy Analytics</title>
    <style>
        :root {{ color-scheme: light; }}
        body {{
            font: 14px Segoe UI, Arial, sans-serif;
            background: #f3f5f8;
            color: #1f2937;
            margin: 0;
        }}
        main {{
            max-width: 1450px;
            margin: auto;
            padding: 24px;
        }}
        header {{
            display: flex;
            justify-content: space-between;
            gap: 16px;
            flex-wrap: wrap;
            align-items: end;
            margin-bottom: 18px;
        }}
        h1 {{ margin: 0 0 4px; }}
        .ctrl, .cards {{
            display: flex;
            gap: 12px;
            flex-wrap: wrap;
            align-items: end;
        }}
        .ctrl {{
            background: white;
            padding: 14px;
            border-radius: 12px;
            box-shadow: 0 2px 10px #0001;
            margin-bottom: 14px;
        }}
        .card, table {{
            background: white;
            box-shadow: 0 2px 10px #0001;
        }}
        .card {{
            padding: 14px;
            border-radius: 12px;
            min-width: 145px;
            flex: 1 1 145px;
        }}
        .v {{
            font-size: 23px;
            font-weight: 700;
            margin-top: 3px;
        }}
        table {{
            width: 100%;
            border-collapse: collapse;
            margin-top: 16px;
            border-radius: 12px;
            overflow: hidden;
        }}
        th, td {{
            padding: 11px;
            border-bottom: 1px solid #eee;
            text-align: left;
            white-space: nowrap;
        }}
        th {{ background: #fafafa; }}
        .n {{ text-align: right; }}
        select, input {{
            padding: 8px;
            border: 1px solid #cfd6df;
            border-radius: 7px;
            background: white;
        }}
        #q {{ min-width: 220px; }}
        small, .muted {{ color: #667085; }}
        .warn {{ color: #9a6700; }}
        .bad {{ color: #b42318; }}
        @media (max-width: 800px) {{
            main {{ padding: 12px; }}
            table {{ display: block; overflow-x: auto; }}
        }}
    </style>
</head>
<body>
    <main>
        <header>
            <div>
                <h1>Scan Proxy Analytics</h1>
                <small>
                    Scanner attribution requires an unambiguous SMB client IP &middot;
                    Reverse DNS via nslookup (optional) &middot;
                    Scan Proxy {escape(VERSION)}
                </small>
            </div>
            <small>Report refresh: {interval_label}</small>
        </header>

        <div class="ctrl">
            <label>
                Period<br>
                <select id="m">
                    <option value="day">Day</option>
                    <option value="week">Week</option>
                    <option value="month">Month</option>
                    <option value="year">Year</option>
                </select>
            </label>
            <label>
                Date<br>
                <input id="d" type="date" value="{latest}">
            </label>
            <label>
                Filter by scanner IP / hostname<br>
                <input id="q" type="search" placeholder="e.g. 192.0.2.10 or scanner-01">
            </label>
            <b id="r"></b>
        </div>

        <div class="cards">
            <div class="card">Scanner<div class="v" id="c">0</div></div>
            <div class="card">Total files<div class="v" id="f">0</div></div>
            <div class="card">Pages<div class="v" id="p">0</div></div>
            <div class="card">Data<div class="v" id="b">0 MB</div></div>
            <div class="card">IP pending<div class="v warn" id="pn">0</div></div>
            <div class="card">IP unresolved<div class="v bad" id="u">0</div></div>
            <div class="card">First scan<div class="v" id="e">-</div></div>
            <div class="card">Last scan<div class="v" id="l">-</div></div>
        </div>

        <table>
            <thead>
                <tr>
                    <th>Scanner IP</th>
                    <th>Hostname</th>
                    <th class="n">Files</th>
                    <th class="n">Pages</th>
                    <th class="n">Avg. pages/file</th>
                    <th class="n">Data MB</th>
                    <th>Last scan</th>
                </tr>
            </thead>
            <tbody id="t"></tbody>
        </table>

        <p class="muted">
            The table contains only successfully transferred files with unambiguously identified
            scanner IPs. Pending and permanently unresolved attributions are shown separately
            and are never assigned to a guessed client or routing identifier.
        </p>
    </main>

    <script>
        const P = {embedded};
        const D = P.data;
        const Q = P.quality;
        const m = document.getElementById('m');
        const d = document.getElementById('d');
        const q = document.getElementById('q');
        const t = document.getElementById('t');
        const pad = n => String(n).padStart(2, '0');
        const safe = value => String(value ?? '').replace(/[&<>"']/g, ch => ({{
            '&': '&amp;', '<': '&lt;', '>': '&gt;', '"': '&quot;', "'": '&#39;'
        }}[ch]));
        const iso = x => `${{x.getFullYear()}}-${{pad(x.getMonth() + 1)}}-${{pad(x.getDate())}}`;
        const pd = s => {{
            const a = s.split('-').map(Number);
            return new Date(a[0], a[1] - 1, a[2]);
        }};
        const fd = s => {{
            const x = pd(s);
            return `${{pad(x.getDate())}}.${{pad(x.getMonth() + 1)}}.${{x.getFullYear()}}`;
        }};
        const dt = s => s ? new Date(s).toLocaleString(
            'en-GB', {{dateStyle: 'short', timeStyle: 'short'}}
        ) : '-';
        const num = x => Number(x || 0).toLocaleString('en-GB');
        const mb = x => (Number(x || 0) / 1048576).toLocaleString(
            'en-GB', {{minimumFractionDigits: 2, maximumFractionDigits: 2}}
        );

        function range() {{
            const x = pd(d.value);
            let a = new Date(x);
            let z = new Date(x);
            if (m.value === 'week') {{
                a.setDate(x.getDate() - ((x.getDay() + 6) % 7));
                z = new Date(a);
                z.setDate(a.getDate() + 6);
            }}
            if (m.value === 'month') {{
                a = new Date(x.getFullYear(), x.getMonth(), 1);
                z = new Date(x.getFullYear(), x.getMonth() + 1, 0);
            }}
            if (m.value === 'year') {{
                a = new Date(x.getFullYear(), 0, 1);
                z = new Date(x.getFullYear(), 11, 31);
            }}
            return [iso(a), iso(z)];
        }}

        function render() {{
            const [A, Z] = range();
            const rows = D.filter(v => v.date >= A && v.date <= Z);
            const quality = Q.filter(v => v.date >= A && v.date <= Z);
            const scanners = new Map();
            let files = 0, pages = 0, bytes = 0, pending = 0, unresolved = 0;
            let first = '', last = '';

            for (const v of quality) {{
                files += v.total;
                pages += v.pages;
                bytes += v.bytes;
                pending += v.pending;
                unresolved += v.unresolved;
                first = !first || (v.first && v.first < first) ? v.first : first;
                last = !last || (v.last && v.last > last) ? v.last : last;
            }}

            for (const v of rows) {{
                const o = scanners.get(v.scanner) || {{
                    hostname: v.hostname || '', f: 0, p: 0, b: 0, l: ''
                }};
                if (!o.hostname && v.hostname) o.hostname = v.hostname;
                o.f += v.files;
                o.p += v.pages;
                o.b += v.bytes;
                o.l = !o.l || v.last > o.l ? v.last : o.l;
                scanners.set(v.scanner, o);
            }}

            document.getElementById('r').textContent = `${{fd(A)}} to ${{fd(Z)}}`;
            document.getElementById('c').textContent = num(scanners.size);
            document.getElementById('f').textContent = num(files);
            document.getElementById('p').textContent = num(pages);
            document.getElementById('b').textContent = mb(bytes) + ' MB';
            document.getElementById('pn').textContent = num(pending);
            document.getElementById('u').textContent = num(unresolved);
            document.getElementById('e').textContent = dt(first);
            document.getElementById('l').textContent = dt(last);

            const filter = q.value.trim().toLowerCase();
            const visible = [...scanners]
                .filter(([ip, o]) => !filter || ip.toLowerCase().includes(filter) ||
                    (o.hostname || '').toLowerCase().includes(filter))
                .sort((a, b) => a[0].localeCompare(b[0], undefined, {{numeric: true}}));

            t.innerHTML = visible.map(([ip, o]) => `
                <tr>
                    <td>${{safe(ip)}}</td>
                    <td>${{o.hostname ? safe(o.hostname) : '<span class="muted">-</span>'}}</td>
                    <td class="n">${{num(o.f)}}</td>
                    <td class="n">${{num(o.p)}}</td>
                    <td class="n">${{o.f ? (o.p / o.f).toLocaleString('en-GB', {{maximumFractionDigits: 1}}) : '-'}}</td>
                    <td class="n">${{mb(o.b)}}</td>
                    <td>${{dt(o.l)}}</td>
                </tr>
            `).join('') || '<tr><td colspan="7">No data</td></tr>';
        }}

        m.onchange = render;
        d.onchange = render;
        q.oninput = render;
        render();
    </script>
</body>
</html>
"""


    def refresh_reports(self, force=False):
        """Refresh only affected CSV periods; HTML keeps the complete history."""
        today = date.today()
        maintenance_due = self._last_maintenance_day != today
        if not force and not self._reports_dirty and not maintenance_due:
            return 0

        full_refresh = bool(force or maintenance_due)
        dirty_dates = set(self._dirty_dates)
        data = self.data()
        scanner_ips = [row["scanner_ip"] for row in data]
        changed_ips = self.refresh_hostnames(scanner_ips)

        if changed_ips and not full_refresh:
            placeholders = ",".join("?" for _ in changed_ips)
            with self.connect() as db:
                rows = db.execute(
                    f"SELECT DISTINCT completed_date FROM scans WHERE scanner_ip IN ({placeholders})",
                    tuple(sorted(changed_ips)),
                ).fetchall()
            dirty_dates.update(row["completed_date"] for row in rows)

        hostnames = self.hostnames()
        quality = self.quality_data()
        cutoff = today - timedelta(days=DAILY_RETENTION_DAYS)

        def periods_for(day):
            week_start = day - timedelta(days=day.weekday())
            iso_week = week_start.isocalendar()
            month_start = date(day.year, day.month, 1)
            next_month = date(day.year + (day.month == 12), (day.month % 12) + 1, 1)
            return (
                ("week", f"{iso_week.year}-W{iso_week.week:02d}-scan-proxy.csv", week_start, week_start + timedelta(days=6)),
                ("month", f"{day.year:04d}-{day.month:02d}-scan-proxy.csv", month_start, next_month - timedelta(days=1)),
                ("year", f"{day.year:04d}-scan-proxy.csv", date(day.year, 1, 1), date(day.year, 12, 31)),
            )

        affected_periods = {"week": set(), "month": set(), "year": set()}
        if not full_refresh:
            for value in dirty_dates:
                try:
                    day = date.fromisoformat(value)
                except ValueError:
                    continue
                for report_type, name, start, finish in periods_for(day):
                    if finish < today:
                        affected_periods[report_type].add((name, start, finish))

        by_day = {}
        groups = {"week": {}, "month": {}, "year": {}}
        for row in data:
            completed_date = date.fromisoformat(row["completed_date"])
            day_key = completed_date.isoformat()
            if full_refresh or day_key in dirty_dates:
                by_day.setdefault(completed_date, []).append(row)

            item = {
                "scanner": row["scanner_ip"],
                "files": row["files"],
                "pages": row["pages"],
                "bytes": row["bytes"],
                "last": row["last_scan"],
            }
            for report_type, name, start, finish in periods_for(completed_date):
                if finish >= today:
                    continue
                period_key = (name, start, finish)
                if not full_refresh and period_key not in affected_periods[report_type]:
                    continue
                period = groups[report_type].setdefault(period_key, {})
                bucket = period.setdefault(
                    item["scanner"],
                    {"files": 0, "pages": 0, "bytes": 0, "last": ""},
                )
                bucket["files"] += item["files"]
                bucket["pages"] += item["pages"]
                bucket["bytes"] += item["bytes"]
                bucket["last"] = max(bucket["last"], item["last"])

        changed = 0
        for completed_date, same_day in by_day.items():
            if completed_date < cutoff:
                continue
            rows = [
                {
                    "Date": completed_date.strftime("%Y-%m-%d"),
                    "Scanner IP": row["scanner_ip"],
                    "Hostname": hostnames.get(row["scanner_ip"], ""),
                    "Files": row["files"],
                    "Pages": row["pages"],
                    "Data_MB": mb(row["bytes"]),
                    "Last_Scan": str(row["last_scan"])[11:16],
                }
                for row in same_day
            ]
            changed += self.write_csv(
                DAILY / f"{completed_date.strftime('%Y-%m-%d')}-scan-proxy.csv",
                ("Date", "Scanner IP", "Hostname", "Files", "Pages", "Data_MB", "Last_Scan"),
                rows,
            )

        for report_type, folder in (("week", WEEKLY), ("month", MONTHLY), ("year", YEARLY)):
            for (name, start, finish), scanners in groups[report_type].items():
                rows = [
                    {
                        "From": start.strftime("%Y-%m-%d"),
                        "To": finish.strftime("%Y-%m-%d"),
                        "Scanner IP": scanner,
                        "Hostname": hostnames.get(scanner, ""),
                        "Files": values["files"],
                        "Pages": values["pages"],
                        "Data_MB": mb(values["bytes"]),
                        "Last_Scan": str(values["last"])[0:16].replace("T", " "),
                    }
                    for scanner, values in sorted(scanners.items())
                ]
                changed += self.write_csv(
                    folder / name,
                    ("From", "To", "Scanner IP", "Hostname", "Files", "Pages", "Data_MB", "Last_Scan"),
                    rows,
                )

        if full_refresh:
            for report_file in DAILY.glob("*-scan-proxy.csv"):
                try:
                    report_date = datetime.strptime(report_file.name[:10], "%Y-%m-%d").date()
                    if report_date < cutoff:
                        report_file.unlink()
                        changed += 1
                        self.log.info("CLEANUP | Deleted expired daily report")
                except ValueError:
                    pass

            week_start = today - timedelta(days=today.weekday())
            iso_week = week_start.isocalendar()
            incomplete = (
                WEEKLY / f"{iso_week.year}-W{iso_week.week:02d}-scan-proxy.csv",
                MONTHLY / f"{today.year:04d}-{today.month:02d}-scan-proxy.csv",
                YEARLY / f"{today.year:04d}-scan-proxy.csv",
            )
            for report_file in incomplete:
                if report_file.exists():
                    report_file.unlink()
                    changed += 1
                    self.log.info(
                        "CLEANUP | Removed incomplete period report | File=%s",
                        report_file,
                    )

        embedded_data = [
            {
                "date": row["completed_date"],
                "scanner": row["scanner_ip"],
                "hostname": hostnames.get(row["scanner_ip"], ""),
                "files": row["files"],
                "pages": row["pages"],
                "bytes": row["bytes"],
                "first": row["first_scan"],
                "last": row["last_scan"],
            }
            for row in data
        ]
        embedded_quality = [
            {
                "date": row["completed_date"],
                "total": int(row["total"] or 0),
                "pages": int(row["pages"] or 0),
                "bytes": int(row["bytes"] or 0),
                "first": row["first_scan"] or "",
                "last": row["last_scan"] or "",
                "resolved": int(row["resolved"] or 0),
                "pending": int(row["pending"] or 0),
                "unresolved": int(row["unresolved"] or 0),
            }
            for row in quality
        ]
        embedded = json.dumps(
            {"data": embedded_data, "quality": embedded_quality},
            ensure_ascii=True,
            separators=(",", ":"),
        )
        # Prevent closing the inline script tag via a malicious DNS PTR name.
        embedded = embedded.replace("<", "\\u003c").replace(">", "\\u003e").replace("&", "\\u0026")
        latest = quality[-1]["completed_date"] if quality else today.isoformat()
        interval_label = (
            f"{REPORT_REFRESH_INTERVAL_SECONDS // 60} minutes"
            if REPORT_REFRESH_INTERVAL_SECONDS % 60 == 0
            else f"{REPORT_REFRESH_INTERVAL_SECONDS} seconds"
        )
        html = self.build_report_html(embedded, latest, interval_label)
        changed += int(self.write_text(REPORT, html))

        if changed:
            self.log.info(
                "ANALYTICS UPDATED | ChangedFiles=%s | Full=%s | DirtyDates=%s",
                changed, full_refresh, len(dirty_dates),
            )
        self._reports_dirty = False
        self._dirty_dates.clear()
        self._last_maintenance_day = today
        return changed


def destination(target, relative, action):
    """Build the destination path and generate unique names for move collisions."""
    dst = target / relative
    ensure_target_dir(dst.parent)
    if action == "copy":
        return dst
    if not dst.exists():
        return dst
    stamp = time.strftime("%Y%m%d_%H%M%S")
    number = 1
    while True:
        suffix = "" if number == 1 else f"_{number}"
        candidate = dst.with_name(f"{dst.stem}_{stamp}{suffix}{dst.suffix}")
        if not candidate.exists():
            return candidate
        number += 1


def transfer(src, target, route_relative, action, reason, label, smb, analytics, log, initial_stat=None):
    """Transfer safely; avoid destination I/O until the source is stable."""
    tmp = None
    dst = None
    try:
        before = initial_stat or src.stat()
        if analytics.pending_move(src):
            return False
        try:
            smb_relative = str(src.relative_to(ROOT))
        except ValueError:
            smb_relative = src.name
        mtime = datetime.fromtimestamp(before.st_mtime, tz=timezone.utc).isoformat().replace("+00:00", "Z")

        resolution = smb.resolve(src, smb_relative, mtime, utc_now())
        if resolution.get("is_open"):
            return False
        if time.time() - before.st_mtime < MIN_AGE:
            return False

        current = src.stat()
        if (before.st_size, before.st_mtime_ns) != (current.st_size, current.st_mtime_ns):
            return False

        dst = destination(target, route_relative, action)
        if norm_path_key(dst) == norm_path_key(src):
            log.error("TRANSFER SKIPPED | Source and destination are identical")
            return False
        if action == "copy" and dst.exists():
            return False

        page_count = pages(src, analytics.log)
        tmp = dst.with_name(dst.name + "." + uuid.uuid4().hex + ".part")
        shutil.copy2(src, tmp)
        after = src.stat()
        if (before.st_size, before.st_mtime_ns) != (after.st_size, after.st_mtime_ns):
            tmp.unlink(missing_ok=True)
            return False
        os.replace(tmp, dst)
        committed = utc_now()

        if action == "move":
            try:
                src.unlink()
            except OSError as exc:
                try:
                    analytics.add_move(src, dst, smb_relative, label, before, mtime, resolution, exc)
                except Exception:
                    try:
                        dst.unlink()
                    except OSError:
                        log.error("CRITICAL | Could not roll back destination after failed move")
                return False

        analytics.add_scan(
            src, smb_relative, dst, label, mtime, resolution,
            completed=committed, page_count=page_count, size_bytes=before.st_size,
        )
        log.info("TRANSFER COMPLETE | Action=%s | Rule=%s", action, reason)
        smb.forget(src)
        return True
    except FileNotFoundError:
        return False
    except Exception as exc:
        log.error("TRANSFER ERROR | ErrorType=%s", type(exc).__name__)
        return False
    finally:
        if tmp is not None:
            tmp.unlink(missing_ok=True)


def process_rules(groups, smb, analytics, log):
    """Read each source once and transfer matching files sequentially."""
    started = time.perf_counter()
    tasks = []
    scanned = 0
    source_errors = set()

    for group in groups:
        if STOP.is_set():
            break
        now = time.monotonic()
        if group.error_until > now:
            continue
        recursive = bool(group.subfolders)
        try:
            candidates = scan_source(group.source, recursive)
            group.error_until = 0.0
        except OSError as exc:
            group.error_until = now + 30.0
            key = norm_path_key(group.source)
            if key not in source_errors:
                source_errors.add(key)
                log.error(
                    "RULE ERROR | Cannot read source directory | Retry=30s | ErrorType=%s",
                    type(exc).__name__,
                )
            continue
        scanned += len(candidates)

        for candidate in candidates:
            if STOP.is_set():
                break
            try:
                relative = candidate.path.relative_to(group.source)
            except ValueError:
                continue

            if group.identifiers and len(relative.parts) == 1:
                stem = candidate.path.stem.lower()
                name = candidate.path.name.lower()
                for key, matching_rules in group.identifiers.items():
                    if stem == key or name.startswith(key + "-"):
                        for rule in matching_rules:
                            tasks.append((candidate, rule, Path(candidate.path.name)))

            for rule in group.subfolders:
                tasks.append((candidate, rule, relative))

    force_snapshot = bool(tasks) and any(
        candidate.stat.st_mtime >= smb.snapshot_wall_time
        for candidate, _, _ in tasks
    )
    snapshot_ms = smb.refresh_open_snapshot(force=force_snapshot) if tasks else 0.0
    transfers = 0
    for candidate, rule, relative in tasks:
        if STOP.is_set():
            break
        # Routing identifiers are not persisted in analytics.
        label = ""
        if transfer(
            candidate.path,
            rule.target,
            relative,
            rule.action,
            f"Excel row {rule.row}",
            label,
            smb,
            analytics,
            log,
            initial_stat=candidate.stat,
        ):
            transfers += 1

    return {
        "ms": (time.perf_counter() - started) * 1000.0,
        "snapshot_ms": snapshot_ms,
        "sources": len(groups),
        "scanned": scanned,
        "matched": len(tasks),
        "transfers": transfers,
    }





def single_instance(log):
    if os.name != "nt":
        return True
    kernel32 = ctypes.WinDLL(
        "kernel32",
        use_last_error=True,
    )
    kernel32.CreateMutexW.restype = ctypes.c_void_p
    kernel32.CreateMutexW.argtypes = (
        ctypes.c_void_p,
        ctypes.c_bool,
        ctypes.c_wchar_p,
    )
    kernel32.CloseHandle.argtypes = (
        ctypes.c_void_p,
    )
    ctypes.set_last_error(0)
    handle = kernel32.CreateMutexW(
        None,
        False,
        "Global\\TrachtiScanProxy",
    )
    already_running = (
        ctypes.get_last_error() == 183
    )
    if not handle or already_running:
        if handle:
            kernel32.CloseHandle(handle)
        log.error(
            "START ABORTED | Another scan-proxy instance is running"
        )
        return False
    return handle


def main():
    log, analytics_log = get_logs()
    mutex = single_instance(log)
    if mutex is False:
        return

    def request_stop(signum, _frame):
        log.info("STOP REQUESTED | Signal=%s", signum)
        STOP.set()

    for signal_name in ("SIGINT", "SIGTERM", "SIGBREAK"):
        stop_signal = getattr(signal, signal_name, None)
        if stop_signal is not None:
            try:
                signal.signal(stop_signal, request_stop)
            except (OSError, ValueError):
                pass

    analytics = Analytics(analytics_log)
    smb = SmbResolver(log)
    analytics.init()
    last_config_mtime = None
    groups = []
    sheet = ""
    next_file_scan = 0.0
    next_pending = time.monotonic() + PENDING_SWEEP_INTERVAL_SECONDS
    next_reports = time.monotonic() + REPORT_REFRESH_INTERVAL_SECONDS
    next_perf = time.monotonic() + PERFORMANCE_LOG_INTERVAL_SECONDS
    perf = {"cycles": 0, "scan_ms": 0.0, "snapshot_ms": 0.0, "scanned": 0,
            "matched": 0, "transfers": 0, "pending_ms": 0.0, "report_ms": 0.0}

    log.info(
        "START | Version=%s | FileScan=%ss | PendingSweep=%ss | "
        "ReportRefresh=%ss | SMBSnapshot=%ss | SMBPending=%smin | 5140Fallback=%ss",
        VERSION, FILE_SCAN_INTERVAL_SECONDS, PENDING_SWEEP_INTERVAL_SECONDS,
        REPORT_REFRESH_INTERVAL_SECONDS, SMB_SNAPSHOT_TTL_SECONDS,
        SMB_EVENT_LOOKBACK_MINUTES, SMB_SHARE_FALLBACK_MAX_AGE_SECONDS,
    )
    analytics_log.info(
        "START | Version=%s | PendingSweep=%ss | ReportRefresh=%ss",
        VERSION, PENDING_SWEEP_INTERVAL_SECONDS, REPORT_REFRESH_INTERVAL_SECONDS,
    )
    print(f"Scanproxy {VERSION} running. Press Ctrl+C to stop.", flush=True)

    try:
        while not STOP.is_set():
            try:
                now = time.monotonic()
                transfers_this_loop = 0

                if now >= next_file_scan:
                    config_mtime = CONFIG_PATH.stat().st_mtime_ns
                    if config_mtime != last_config_mtime:
                        sheet, raw_rules = load_rules()
                        groups, active_count = compile_rules(raw_rules, log)
                        last_config_mtime = config_mtime
                        log.info(
                            "WORKBOOK LOADED | Sheet=%s | Rules=%s | Enabled=%s | Sources=%s",
                            sheet, len(raw_rules), active_count, len(groups),
                        )

                    stats = process_rules(groups, smb, analytics, log)
                    transfers_this_loop = stats["transfers"]
                    perf["cycles"] += 1
                    perf["scan_ms"] += stats["ms"]
                    perf["snapshot_ms"] += stats["snapshot_ms"]
                    perf["scanned"] += stats["scanned"]
                    perf["matched"] += stats["matched"]
                    perf["transfers"] += stats["transfers"]
                    next_file_scan = time.monotonic() + FILE_SCAN_INTERVAL_SECONDS

                if not transfers_this_loop:
                    analytics.retry_moves()

                    if now >= next_pending:
                        t0 = time.perf_counter()
                        analytics.resolve_pending(smb)
                        perf["pending_ms"] += (time.perf_counter() - t0) * 1000.0
                        next_pending = time.monotonic() + PENDING_SWEEP_INTERVAL_SECONDS

                    if now >= next_reports:
                        t0 = time.perf_counter()
                        analytics.refresh_reports()
                        perf["report_ms"] += (time.perf_counter() - t0) * 1000.0
                        next_reports = time.monotonic() + REPORT_REFRESH_INTERVAL_SECONDS

                if now >= next_perf:
                    cycles = max(1, perf["cycles"])
                    log.info(
                        "PERFORMANCE | ScanAvg=%.1fms | SMBSnapshot=%.1fms | "
                        "Sources=%s | Files=%s | Matches=%s | Transfers=%s | "
                        "Pending=%.1fms | Reports=%.1fms",
                        perf["scan_ms"] / cycles,
                        perf["snapshot_ms"] / cycles,
                        len(groups), perf["scanned"], perf["matched"], perf["transfers"],
                        perf["pending_ms"], perf["report_ms"],
                    )
                    perf = {"cycles": 0, "scan_ms": 0.0, "snapshot_ms": 0.0, "scanned": 0,
                            "matched": 0, "transfers": 0, "pending_ms": 0.0, "report_ms": 0.0}
                    next_perf = time.monotonic() + PERFORMANCE_LOG_INTERVAL_SECONDS

            except Exception:
                log.error("PROCESSING LOOP FAILED | ErrorType=%s", sys.exc_info()[0].__name__)
            STOP.wait(MAIN_LOOP_SLEEP_SECONDS)

    finally:
        try:
            analytics.resolve_pending(smb)
            analytics.refresh_reports(force=True)
        except Exception:
            analytics_log.error("ANALYTICS ERROR | Shutdown | ErrorType=%s", sys.exc_info()[0].__name__)
        analytics.close()
        if os.name == "nt" and mutex not in (True, False):
            ctypes.WinDLL("kernel32").CloseHandle(ctypes.c_void_p(mutex))
        analytics_log.info("STOPPED")
        log.info("STOPPED")
        print("Scan Proxy stopped.", flush=True)

if __name__ == "__main__":
    main()