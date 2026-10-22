"""Cross-platform unit and filesystem integration tests for Scan Proxy."""
from __future__ import annotations

import importlib.util
import logging
import os
import sqlite3
import sys
import tempfile
import time
import unittest
from datetime import datetime, timezone
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
MODULE = ROOT / "src" / "scan_proxy.py"
SPEC = importlib.util.spec_from_file_location("scan_proxy", MODULE)
assert SPEC is not None and SPEC.loader is not None
scan_proxy = importlib.util.module_from_spec(SPEC)
sys.modules[SPEC.name] = scan_proxy
SPEC.loader.exec_module(scan_proxy)


class ScanProxyTests(unittest.TestCase):
    def test_version_and_neutral_smb_defaults(self):
        self.assertEqual(scan_proxy.VERSION, "2.13.1")
        self.assertEqual(scan_proxy.SMB_IGNORED_NETWORKS, ())
        # An explicit SMB user is required for production attribution.
        if not os.environ.get("SCAN_PROXY_SMB_USER"):
            self.assertEqual(scan_proxy.SMB_USER, "")

    def test_rule_activation_and_ip_validation(self):
        for value in (True, "true", "1", "yes", "on", "enabled", "x"):
            self.assertTrue(scan_proxy.is_enabled(value))
        for value in (False, "false", "no", "off", "", None):
            self.assertFalse(scan_proxy.is_enabled(value))
        self.assertEqual(scan_proxy.normalize_ip("127.0.0.1"), "")
        self.assertEqual(scan_proxy.normalize_ip("::1"), "")
        self.assertEqual(scan_proxy.normalize_ip("not-an-ip"), "")
        self.assertEqual(scan_proxy.normalize_ip("192.0.2.10"), "192.0.2.10")

    def test_snapshot_off_without_smb_configuration(self):
        logger = logging.getLogger("test_snapshot")
        smb = scan_proxy.SmbResolver(logger)
        with patch.object(scan_proxy, "SMB_USER", ""):
            self.assertGreaterEqual(smb.refresh_open_snapshot(force=True), 0)
            self.assertFalse(smb.resolve("example.pdf")["resolved"])
            self.assertEqual(smb.resolve_batch([{"id": "test"}])["test"]["client"], "")

    def test_snapshot_cache_never_reuses_ip_for_a_different_file_version(self):
        logger = logging.getLogger("test_snapshot_cache")
        smb = scan_proxy.SmbResolver(logger)
        source = Path("synthetic.pdf")
        path_key = scan_proxy.norm_path_key(source)
        smb.snapshot_paths[path_key] = {"192.0.2.22"}
        first = smb.resolve(source, mtime="2026-10-22T08:00:00Z")
        self.assertTrue(first["is_open"])
        smb.snapshot_paths.clear()
        cached = smb.resolve(source, mtime="2026-10-22T08:00:00Z")
        self.assertEqual(cached["client"], "192.0.2.22")
        self.assertFalse(cached["is_open"])
        new_file = smb.resolve(source, mtime="2026-10-22T08:01:00Z")
        self.assertFalse(new_file["resolved"])

    def test_retry_backoff(self):
        self.assertEqual(scan_proxy.pending_retry_delay(0), 30)
        self.assertEqual(scan_proxy.pending_retry_delay(1), 120)
        self.assertEqual(scan_proxy.pending_retry_delay(999), 43200)

    def test_example_workbook_contains_no_live_rules(self):
        path = ROOT / "config" / "scan-proxy-config.example.xlsx"
        workbook = load_workbook(path, read_only=True)
        try:
            sheet = workbook["Rules"]
            self.assertEqual(tuple(x.value for x in sheet[1]),
                             ("Identifier", "Source", "Destination", "Mode", "Enabled", "Action", "Scanner"))
            self.assertTrue(all(row[4] is False for row in sheet.iter_rows(min_row=2, values_only=True)))
        finally:
            workbook.close()

    def test_rules_and_filesystem_routing(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            (root / "incoming").mkdir()
            (root / "incoming" / "DOC-test.txt").write_text("synthetic test data", encoding="utf-8")
            src = root / "incoming" / "DOC-test.txt"
            old_timestamp = time.time() - 10
            os.utime(src, (old_timestamp, old_timestamp))
            config = root / "rules.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Rules"
            sheet.append(["Identifier", "Source", "Destination", "Mode", "Enabled", "Action", "Scanner"])
            sheet.append(["DOC", "incoming", "routed", "identifier", True, "move", "test-only"])
            workbook.save(config)
            locations = {
                "ROOT": root, "CONFIG_PATH": config,
                "ANALYTICS": root / "analytics", "STATE": root / "analytics/_state",
                "DB": root / "analytics/_state/scan-proxy.db",
                "REPORT": root / "analytics/analytics.html",
                "DAILY": root / "analytics/daily", "WEEKLY": root / "analytics/weekly",
                "MONTHLY": root / "analytics/monthly", "YEARLY": root / "analytics/yearly",
            }
            logger = logging.getLogger("test_routing")
            previous = {name: getattr(scan_proxy, name) for name in locations}
            try:
                for name, value in locations.items():
                    setattr(scan_proxy, name, value)
                _, rules = scan_proxy.load_rules()
                groups, active_count = scan_proxy.compile_rules(rules, logger)
                self.assertEqual(active_count, 1)
                self.assertEqual(groups[0].identifiers["doc"][0].action, "move")
                analytics = scan_proxy.Analytics(logger)
                analytics.init()
                with patch.object(scan_proxy, "SMB_USER", ""):
                    result = scan_proxy.process_rules(groups, scan_proxy.SmbResolver(logger), analytics, logger)
                self.assertEqual(result["transfers"], 1)
                self.assertFalse(src.exists())
                self.assertTrue((root / "routed" / "DOC-test.txt").is_file())
                row = analytics.connect().execute(
                    "SELECT resolution_state, original_path, relative_path, completed_path, routing_label FROM scans"
                ).fetchone()
                self.assertEqual(row["resolution_state"], "unresolved")
                self.assertEqual(row["original_path"], "", "Do not retain paths when SMB is off")
                self.assertEqual(row["completed_path"], "", "Completed paths are never stored")
                self.assertEqual(row["routing_label"], "", "Labels are not stored")
                self.assertTrue((root / "analytics/analytics.html").exists())
                html = (root / "analytics/analytics.html").read_text(encoding="utf-8")
                self.assertIn("Scan Proxy Analytics", html)
                self.assertIn("const safe =", html)
                self.assertNotIn("test-only", html)
                analytics.close()
            finally:
                for name, value in previous.items():
                    setattr(scan_proxy, name, value)

    def test_pending_resolution_scrubs_paths(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            location = root / "analytics"
            previous = {name: getattr(scan_proxy, name) for name in
                        ("ANALYTICS", "STATE", "DB", "REPORT", "DAILY", "WEEKLY", "MONTHLY", "YEARLY")}
            locations = {
                "ANALYTICS": location, "STATE": location / "_state",
                "DB": location / "_state/scan-proxy.db", "REPORT": location / "analytics.html",
                "DAILY": location / "daily", "WEEKLY": location / "weekly",
                "MONTHLY": location / "monthly", "YEARLY": location / "yearly",
            }
            logger = logging.getLogger("test_pending")
            try:
                for name, value in locations.items():
                    setattr(scan_proxy, name, value)
                analytics = scan_proxy.Analytics(logger)
                analytics.init()
                file = root / "synthetic.txt"
                file.write_text("sample", encoding="utf-8")
                with patch.object(scan_proxy, "SMB_USER", "scanner-test-account"):
                    analytics.add_scan(
                        file, "synthetic.txt", file, "", datetime.now(timezone.utc).isoformat(),
                        scan_proxy.SmbResolver.empty(), page_count=1, size_bytes=6,
                    )
                analytics.connect().execute("UPDATE scans SET next_resolve_epoch=0")
                analytics.connect().commit()

                class StubSmb:
                    def resolve_batch(self, rows):
                        return {str(row["id"]): {
                            "client": "192.0.2.15", "source": "Security-5145",
                            "reason": "Synthetic event", "confidence": "ExactPathRecent", "record_id": 10,
                        } for row in rows}
                self.assertEqual(analytics.resolve_pending(StubSmb()), 1)
                row = analytics.connect().execute("SELECT * FROM scans").fetchone()
                self.assertEqual(row["resolution_state"], "resolved")
                self.assertEqual(row["scanner_ip"], "192.0.2.15")
                for column in ("original_path", "relative_path", "completed_path", "routing_label"):
                    self.assertEqual(row[column], "", column)
                analytics.close()
            finally:
                for name, value in previous.items():
                    setattr(scan_proxy, name, value)

    def test_move_collision_and_copy_destination(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            src = root / "sample.pdf"
            src.write_bytes(b"existing")
            self.assertEqual(scan_proxy.destination(root, Path(src.name), "copy"), src)
            self.assertNotEqual(scan_proxy.destination(root, Path(src.name), "move"), src)

    def test_powershell_helper_mentions_both_security_events(self):
        helper = (ROOT / "scripts" / "scan-proxy-smb.ps1").read_text(encoding="utf-8")
        self.assertIn("'OpenSnapshot'", helper)
        self.assertIn("'ResolveBatch'", helper)
        self.assertIn("Id=@(5145,5140)", helper)
        self.assertIn("ShareMatches", helper)
        self.assertIn("Trachti", helper)
        self.assertIn("[string]$UserName=''", helper)
        self.assertIn("[string]$IgnoredNetworks=''", helper)


if __name__ == "__main__":
    unittest.main()
