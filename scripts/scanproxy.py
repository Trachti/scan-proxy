from __future__ import annotations

import csv
import ctypes
import ipaddress
import json
import logging
import os
import re
import shutil
import signal
import sqlite3
import subprocess
import sys
import threading
import time
import uuid
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal, ROUND_HALF_UP
from html import escape
from pathlib import Path

from openpyxl import load_workbook

try:
    from pypdf import PdfReader
except ImportError:
    PdfReader = None


VERSION = "2.11.0-public"

# +------------+---------+-------------------------------------------------------------------------+
# | Date       | Author  | Change                                                                  |
# +------------+---------+-------------------------------------------------------------------------+
# | 01.10.2026 | Trachti | Public release: translated user-facing text and comments to English;    |
# |            |         | removed organization-specific defaults from the published configuration.|
# +------------+---------+-------------------------------------------------------------------------+
# | 14.08.2026 | Trachti | Improved readability: documented functions and split long lines while  |
# |            |         | preserving the existing program logic and optimizations.                |
# +------------+---------+-------------------------------------------------------------------------+
# | 14.08.2026 | Trachti | Split logs; configurable SMB interval; reduced file-scan frequency and  |
# |            |         | optimized SQLite/SMB caching to reduce CPU usage.                       |
# +------------+---------+-------------------------------------------------------------------------+
# | 14.08.2026 | Trachti | Compact release; SMB analytics/reports updated every 10 minutes.        |
# +------------+---------+-------------------------------------------------------------------------+


# ==============================================================================
# Paths and files
# ==============================================================================

ROOT = Path(os.environ.get("SCANPROXY_ROOT", r"C:\ScanProxy\scan"))
BASE = Path(
    sys.executable if getattr(sys, "frozen", False) else __file__
).resolve().parent
CONFIG = BASE / "it-scan-config.xlsx"
SMB_HELPER = BASE / "scanproxy-smb.ps1"

# ==============================================================================
# Manually configurable values
# ==============================================================================

# How often rules and source folders are scanned.
FILE_SCAN_INTERVAL_SECONDS = 5

# Interval for SMB re-resolution and analytics refreshes.
SMB_QUERY_INTERVAL_SECONDS = 600  # 10 minutes

# Lookback window for Windows Security Event 5145 during an SMB query.
SMB_EVENT_LOOKBACK_MINUTES = 10

# Networks listed here are ignored as scanner clients. Configure them with
# SCANPROXY_IGNORED_NETWORKS as a comma- or semicolon-separated list.
SMB_IGNORED_NETWORKS = tuple(
    network.strip()
    for network in re.split(r"[,;]", os.environ.get("SCANPROXY_IGNORED_NETWORKS", ""))
    if network.strip()
)

# Windows/SMB account whose sessions and file access events are evaluated.
SMB_USER = os.environ.get("SCANPROXY_SMB_USER", "scan-service")

# ==============================================================================
# Internal paths and runtime values
# ==============================================================================

LOG_ROOT = BASE / "logs"
SCANPROXY_LOG_DIR = LOG_ROOT / "scanproxy"
ANALYTICS_LOG_DIR = LOG_ROOT / "analytics"
ANALYTICS = BASE / "analytics"
STATE = ANALYTICS / "_state"
REPORT = ANALYTICS / "analytics.html"
DB = STATE / "scanproxy.db"
DAILY = ANALYTICS / "daily"
WEEKLY = ANALYTICS / "weekly"
MONTHLY = ANALYTICS / "monthly"
YEARLY = ANALYTICS / "yearly"

# A file must be at least this old before processing to avoid files still being written.
MIN_AGE = 3

# The main loop wakes briefly once per second.
# Expensive file and SMB checks have separate intervals and do not run on every loop.
MAIN_LOOP_SLEEP_SECONDS = 1

# Maximum duration of a single PowerShell SMB call.
SMB_TIMEOUT = 8

# Scanner clients that cannot be resolved immediately may be retried for up to 48 hours.
PENDING_MAX_HOURS = 48

# Cache lifetimes for SMB results. Open files and misses are cached briefly; resolved results longer.
SMB_OPEN_CACHE_SECONDS = 15
SMB_RESOLVED_CACHE_SECONDS = 300
SMB_MISS_CACHE_SECONDS = 30

# Retention period for daily CSV files.
DAILY_RETENTION_DAYS = 90

# Retry interval for incomplete move operations.
MOVE_RETRY_SECONDS = 5

# This cache avoids a SQLite query for every discovered file.
PENDING_CACHE_SECONDS = 5
SMB_IGNORED_NETS = tuple(
    ipaddress.ip_network(network, strict=False)
    for network in SMB_IGNORED_NETWORKS
)
STOP = threading.Event()

# ==============================================================================
# General helper functions
# ==============================================================================

def p(value):
    path = Path(str(value or ".").strip())
    if path.is_absolute():
        return path
    return ROOT / path

# -----------------------------------------------------------------
# Accept common spellings for an enabled value in Excel.
# -----------------------------------------------------------------
def active(value):
    return str(value).strip().lower() in {
        "true",
        "1",
        "yes",
        "x",
    }

def norm_ip(value):
    text = str(value or "").strip().strip("[]")
    try:
        addr = ipaddress.ip_address(text)
        if isinstance(addr, ipaddress.IPv6Address) and addr.ipv4_mapped:
            addr = addr.ipv4_mapped
        if str(addr) in {
            "127.0.0.1",
            "::1",
            "0.0.0.0",
            "::",
        }:
            return ""

        if any(
            addr.version == network.version and addr in network
            for network in SMB_IGNORED_NETS
        ):
            return ""
        return str(addr)
    except ValueError:
        return ""

def utc_now():
    return datetime.now(timezone.utc).isoformat().replace("+00:00", "Z")

def parse_utc(value):
    try:
        return datetime.fromisoformat(
            str(value).replace("Z", "+00:00")
        ).astimezone(timezone.utc)
    except (ValueError, TypeError):
        return None

# ----------------------------
# Convert bytes to megabytes.
# ----------------------------

def mb(size):
    value = Decimal(int(size or 0)) / Decimal(1048576)
    rounded = value.quantize(
        Decimal("0.01"),
        rounding=ROUND_HALF_UP,
    )
    return f"{rounded:.2f}"

# -----------------------------------
# Determine the number of pages in a PDF.
# Non-PDF files count as one page.
# -----------------------------------

def pages(file, log):
    if file.suffix.lower() != ".pdf":
        return 1
    try:
        if PdfReader:
            reader = PdfReader(
                str(file),
                strict=False,
            )
            return max(1, len(reader.pages))
        content = file.read_bytes()
        return max(
            1,
            len(re.findall(rb"/Type\s*/Page\b", content)),
        )
    except Exception:
        log.warning(
            "ANALYTICS WARNING | "
            "PDF page count could not be read | "
            "File=%s",
            file,
        )
        return 1

# ==============================================================================
# Logging
# ==============================================================================

class DailyLog(logging.Handler):
    def __init__(self, folder):
        super().__init__()
        self.folder = folder
        self.day = None
        self.stream = None
        folder.mkdir(
            parents=True,
            exist_ok=True,
        )

# ------------------------------------------------------------------------------------------------------------------
# Write log entries to a daily file and rotate automatically when the date changes.
# ------------------------------------------------------------------------------------------------------------------
    def emit(self, record):
        try:
            day = datetime.now().strftime("%d.%m.%Y")
            if day != self.day:
                if self.stream:
                    self.stream.close()
                log_file = self.folder / f"{day}-scanproxy.log"
                self.stream = log_file.open(
                    "a",
                    encoding="utf-8",
                )
                self.day = day
            self.stream.write(
                self.format(record) + "\n"
            )
            self.stream.flush()
        except Exception:
            self.handleError(record)
    def close(self):
        if self.stream:
            self.stream.close()
        super().close()

def get_logger(name, folder):
    log = logging.getLogger(name)
    log.setLevel(logging.INFO)
    log.propagate = False
    if not log.handlers:
        handler = DailyLog(folder)
        handler.setFormatter(
            logging.Formatter(
                "%(asctime)s | %(levelname)s | %(message)s"
            )
        )
        log.addHandler(handler)
    return log

# --------------------------------------------------------
# Create the ScanProxy and analytics loggers.
# --------------------------------------------------------

def get_logs():
    scanproxy_log = get_logger(
        "scanproxy",
        SCANPROXY_LOG_DIR,
    )
    analytics_log = get_logger(
        "analytics",
        ANALYTICS_LOG_DIR,
    )
    return scanproxy_log, analytics_log

# ==============================================================================
# Excel rules
# ==============================================================================


# -------------------------------------------------------------------
# Load rules from Excel and prefer a worksheet named Rules.
# -------------------------------------------------------------------

def load_rules():
    workbook = load_workbook(
        CONFIG,
        read_only=True,
        data_only=True,
    )
    try:
        worksheets = sorted(
            workbook.worksheets,
            key=lambda sheet: sheet.title.strip().lower() != "rules",
        )
        for worksheet in worksheets:
            header = next(
                worksheet.iter_rows(
                    min_row=1,
                    max_row=1,
                    values_only=True,
                ),
                (),
            )
            columns = {
                str(value).strip().lower(): index
                for index, value in enumerate(header)
                if value is not None
            }
            required = {
                "id",
                "source",
                "target",
                "mode",
            }
            has_active_column = (
                "enabled" in columns
                or "status" in columns
            )
            if required <= columns.keys() and has_active_column:
                break
        else:
            raise ValueError(
                "No worksheet with ID, Source, Target, Mode and Enabled/Status was found"
            )
        columns["enabled"] = columns.get(
            "enabled",
            columns.get("status"),
        )
        def val(row, name, default=None):
            index = columns.get(name)
            if index is not None and index < len(row):
                return row[index]
            return default
        rules = []
        for row_number, row in enumerate(
            worksheet.iter_rows(
                min_row=2,
                values_only=True,
            ),
            2,
        ):
            if not any(value not in (None, "") for value in row):
                continue
            rules.append(
                (
                    row_number,
                    val(row, "id"),
                    val(row, "source"),
                    val(row, "target"),
                    val(row, "mode"),
                    val(row, "enabled"),
                    val(row, "action", "move") or "move",
                    val(row, "scanner"),
                )
            )
        return worksheet.title, rules
    finally:
        workbook.close()

# ==============================================================================
# SMB resolution
# ==============================================================================


class Smb:
    def __init__(self, log):
        self.log = log
        self.cache = {}
    def resolve(
        self,
        file,
        relative="",
        mtime="",
        reference="",
    ):

# ---------------------------------------------------------------
# Call scanproxy-smb.ps1 for a file and return the scanner IP address.
# ---------------------------------------------------------------

        key = (
            str(file).lower(),
            str(mtime),
        )
        now = time.monotonic()
        cached = self.cache.get(key)
        if cached and cached[1] > now:
            return dict(cached[0])
        empty = {
            "client": "",
            "resolved": False,
            "source": "None",
            "reason": "No unambiguous SMB client IP mapping",
            "record_id": 0,
            "is_open": False,
            "confidence": "None",
        }
        if os.name != "nt" or not SMB_HELPER.is_file():
            return empty
        if not relative:
            try:
                relative = str(file.relative_to(ROOT))
            except ValueError:
                relative = file.name
        command = [
            "powershell.exe",
            "-NoLogo",
            "-NoProfile",
            "-NonInteractive",
            "-ExecutionPolicy",
            "Bypass",
            "-File",
            str(SMB_HELPER),
            "-Mode",
            "Resolve",
            "-FilePath",
            str(file),
            "-RelativePath",
            relative,
            "-UserName",
            SMB_USER,
            "-LookbackMinutes",
            str(SMB_EVENT_LOOKBACK_MINUTES),
            "-ReferenceTimeUtc",
            reference or utc_now(),
            "-IgnoredNetworks",
            ";".join(SMB_IGNORED_NETWORKS),
        ]
        try:
            completed_process = subprocess.run(
                command,
                capture_output=True,
                text=True,
                encoding="utf-8",
                errors="replace",
                timeout=SMB_TIMEOUT,
                creationflags=getattr(
                    subprocess,
                    "CREATE_NO_WINDOW",
                    0,
                ),
            )
            if completed_process.returncode:
                error_text = (
                    completed_process.stderr
                    or f"ExitCode={completed_process.returncode}"
                ).strip()
                raise RuntimeError(error_text)
            data = next(
                (
                    json.loads(line.lstrip("\ufeff"))
                    for line in reversed(
                        completed_process.stdout.splitlines()
                    )
                    if line.strip().startswith("{")
                ),
                {},
            )
            client = norm_ip(
                data.get("Client")
            )
            result = {
                "client": client,
                "resolved": bool(client),
                "source": str(
                    data.get("Source") or "None"
                ),
                "reason": str(
                    data.get("Reason") or empty["reason"]
                ),
                "record_id": int(
                    data.get("RecordId") or 0
                ),
                "is_open": bool(
                    client and data.get("Open")
                ),
                "confidence": str(
                    data.get("Confidence") or "None"
                ),
            }
        except Exception as exc:
            self.log.info(
                "SMB CLIENT NOT RESOLVED | "
                "File=%s | Reason=%s",
                file,
                exc,
            )
            result = empty
        if result["is_open"]:
            ttl = SMB_OPEN_CACHE_SECONDS
        elif result["resolved"]:
            ttl = SMB_RESOLVED_CACHE_SECONDS
        else:
            ttl = SMB_MISS_CACHE_SECONDS

        self.cache[key] = (
            dict(result),
            now + ttl,
        )
# Remove expired cache entries in batches when many files have been seen.
        if len(self.cache) > 2000:
            self.cache = {
                cache_key: cache_value
                for cache_key, cache_value in self.cache.items()
                if cache_value[1] > now
            }
        return result


# ==============================================================================
# Analytics and database
# ==============================================================================


class Analytics:

# ---------------------------------------------------------------------
# Initialize logging, the retry timer, and the pending-move cache.
# ---------------------------------------------------------------------

    def __init__(self, log):
        self.log = log
        self.last_move_retry = 0.0
        self._pending_cache = set()
        self._pending_cache_until = 0.0

# ----------------------------------------------------------------
# Open a SQLite connection and return rows as sqlite3.Row objects.
# ----------------------------------------------------------------

    def connect(self):
        db = sqlite3.connect(
            DB,
            timeout=10,
        )
        db.row_factory = sqlite3.Row
        return db

# --------------------------------------------------
# Create analytics folders, tables, and indexes.
# Enable SQLite WAL mode and synchronous=NORMAL.
# --------------------------------------------------

    def init(self):
        for folder in (
            ANALYTICS,
            STATE,
            DAILY,
            WEEKLY,
            MONTHLY,
            YEARLY,
        ):
            folder.mkdir(
                parents=True,
                exist_ok=True,
            )
        with self.connect() as db:
            db.execute(
                "PRAGMA journal_mode=WAL"
            )
            db.execute(
                "PRAGMA synchronous=NORMAL"
            )
            db.executescript(
                """
                CREATE TABLE IF NOT EXISTS scans (
                    id TEXT PRIMARY KEY,
                    completed_at_utc TEXT NOT NULL,
                    completed_at_local TEXT NOT NULL,
                    completed_date TEXT NOT NULL,
                    original_path TEXT NOT NULL,
                    relative_path TEXT NOT NULL,
                    completed_path TEXT NOT NULL,
                    routing_label TEXT NOT NULL,
                    file_mtime_utc TEXT NOT NULL,
                    pages INTEGER NOT NULL,
                    size_bytes INTEGER NOT NULL,
                    scanner_ip TEXT,
                    resolution_source TEXT NOT NULL,
                    resolution_reason TEXT NOT NULL,
                    resolution_confidence TEXT NOT NULL,
                    resolution_record_id INTEGER NOT NULL DEFAULT 0,
                    resolution_state TEXT NOT NULL,
                    resolve_attempts INTEGER NOT NULL DEFAULT 0,
                    next_resolve_epoch REAL NOT NULL DEFAULT 0,
                    resolve_deadline_utc TEXT NOT NULL,
                    created_at_utc TEXT NOT NULL,
                    updated_at_utc TEXT NOT NULL
                );

                CREATE INDEX IF NOT EXISTS idx_scans_pending
                    ON scans (
                        resolution_state,
                        next_resolve_epoch
                    );

                CREATE INDEX IF NOT EXISTS idx_scans_report
                    ON scans (
                        completed_date,
                        scanner_ip,
                        resolution_state
                    );

                CREATE TABLE IF NOT EXISTS pending_moves (
                    id TEXT PRIMARY KEY,
                    source_norm TEXT NOT NULL,
                    source_path TEXT NOT NULL,
                    target_path TEXT NOT NULL,
                    relative_path TEXT NOT NULL,
                    routing_label TEXT NOT NULL,
                    source_size INTEGER NOT NULL,
                    source_mtime_ns INTEGER NOT NULL,
                    file_mtime_utc TEXT NOT NULL,
                    committed_at_utc TEXT NOT NULL,
                    resolution_json TEXT NOT NULL,
                    reason TEXT NOT NULL,
                    state TEXT NOT NULL,
                    created_at_utc TEXT NOT NULL,
                    last_attempt_utc TEXT NOT NULL
                );

                CREATE INDEX IF NOT EXISTS idx_pending_moves_source
                    ON pending_moves (
                        source_norm,
                        state
                    );
                """
            )
        self.refresh_reports()

# ----------------------------------------------------------------------------------------------
# Check whether SQLite contains a pending or blocked move for a source file.
# The list is cached in memory between checks.
# ----------------------------------------------------------------------------------------------

    def pending_move(self, source):
        now = time.monotonic()
        if now >= self._pending_cache_until:
            with self.connect() as db:
                rows = db.execute(
                    """
                    SELECT source_norm
                    FROM pending_moves
                    WHERE state IN ('pending', 'blocked')
                    """
                ).fetchall()
            self._pending_cache = {
                row["source_norm"]
                for row in rows
            }
            self._pending_cache_until = (
                now + PENDING_CACHE_SECONDS
            )
        key = os.path.normcase(
            os.path.normpath(
                str(source)
            )
        )
        return key in self._pending_cache
    def add_move(
        self,
        source,
        target,
        relative,
        label,
        stat,
        mtime,
        resolution,
        reason,
    ):
        
# --------------------------------------------------------------------------------------------------------------------------
# Store a move whose destination was written but whose source could not yet be deleted.
# --------------------------------------------------------------------------------------------------------------------------

        now = utc_now()
        move_id = uuid.uuid4().hex
        source_norm = os.path.normcase(
            os.path.normpath(
                str(source)
            )
        )
        values = (
            move_id,
            source_norm,
            str(source),
            str(target),
            str(relative),
            label,
            int(stat.st_size),
            int(stat.st_mtime_ns),
            mtime,
            now,
            json.dumps(resolution),
            str(reason),
            "pending",
            now,
            now,
        )
        with self.connect() as db:
            db.execute(
                """
                INSERT INTO pending_moves
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                values,
            )
        self._pending_cache.add(source_norm)
    def add_scan(
        self,
        source,
        relative,
        target,
        label,
        mtime,
        resolution,
        scan_id=None,
        completed=None,
    ):
        
# -----------------------------------------------------------------------
# Store a successfully completed scan in SQLite.
# If no scanner IP is known, store the record as pending.
# -----------------------------------------------------------------------

        done = (
            parse_utc(completed)
            or datetime.now(timezone.utc)
        )
        local = done.astimezone()
        now = utc_now()

        client = norm_ip(
            (resolution or {}).get("client")
        )
        state = (
            "resolved"
            if client
            else "pending"
        )
        deadline = (
            datetime.now(timezone.utc)
            + timedelta(hours=PENDING_MAX_HOURS)
        ).isoformat().replace(
            "+00:00",
            "Z",
        )
        values = (
            scan_id or uuid.uuid4().hex,
            done.isoformat().replace("+00:00", "Z"),
            local.isoformat(timespec="seconds"),
            local.date().isoformat(),
            str(source),
            str(relative),
            str(target),
            str(label or ""),
            mtime,
            pages(target, self.log),
            target.stat().st_size,
            client or None,
            str(
                (resolution or {}).get("source")
                or "None"
            ),
            str(
                (resolution or {}).get("reason")
                or ""
            ),
            str(
                (resolution or {}).get("confidence")
                or "None"
            ),
            int(
                (resolution or {}).get("record_id")
                or 0
            ),
            state,
            0,
            time.time() + SMB_QUERY_INTERVAL_SECONDS,
            deadline,
            now,
            now,
        )
        with self.connect() as db:
            db.execute(
                """
                INSERT OR IGNORE INTO scans
                VALUES (
                    ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?,
                    ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?
                )
                """,
                values,
            )
        self.log.info(
            "ANALYTICS | Scanner=%s | File=%s | Status=%s",
            client or "pending",
            target,
            state,
        )

# -------------------------------------------------------------
# Periodically retry incomplete move operations.
# -------------------------------------------------------------

    def retry_moves(self):
        now = time.monotonic()
        if now - self.last_move_retry < MOVE_RETRY_SECONDS:
            return
        self.last_move_retry = now
        
# --------------------------------------------------------------------
# Invalidate the cache during retries because the database content changes.
# --------------------------------------------------------------------

        self._pending_cache_until = 0.0
        with self.connect() as db:
            rows = db.execute(
                """
                SELECT *
                FROM pending_moves
                WHERE state = 'pending'
                ORDER BY created_at_utc
                LIMIT 10
                """
            ).fetchall()
        for row in rows:
            src = Path(row["source_path"])
            dst = Path(row["target_path"])
            if not dst.exists():
                with self.connect() as db:
                    db.execute(
                        "DELETE FROM pending_moves WHERE id = ?",
                        (row["id"],),
                    )
                continue
            try:
                stat = src.stat()
            except FileNotFoundError:
                stat = None
            except OSError:
                continue
            source_changed = (
                stat
                and (
                    stat.st_size != row["source_size"]
                    or stat.st_mtime_ns != row["source_mtime_ns"]
                )
            )
            if source_changed:
                try:
                    dst.unlink()
                    state = None
                except OSError:
                    state = "blocked"
                with self.connect() as db:
                    if state:
                        db.execute(
                            """
                            UPDATE pending_moves
                            SET state = 'blocked',
                                last_attempt_utc = ?
                            WHERE id = ?
                            """,
                            (
                                utc_now(),
                                row["id"],
                            ),
                        )
                    else:
                        db.execute(
                            "DELETE FROM pending_moves WHERE id = ?",
                            (row["id"],),
                        )
                continue
            if stat:
                try:
                    src.unlink()
                except OSError:
                    continue
            try:
                resolution = json.loads(
                    row["resolution_json"] or "{}"
                )
            except json.JSONDecodeError:
                resolution = {}
            self.add_scan(
                src,
                row["relative_path"],
                dst,
                row["routing_label"],
                row["file_mtime_utc"],
                resolution,
                row["id"],
                row["committed_at_utc"],
            )
            with self.connect() as db:
                db.execute(
                    "DELETE FROM pending_moves WHERE id = ?",
                    (row["id"],),
                )

# ------------------------------------------------------------------------------------------------------------------------------
# Retry SMB resolution for pending analytics records and mark expired records as unresolved.
# ------------------------------------------------------------------------------------------------------------------------------

    def resolve_pending(self, smb):
        now = time.time()
        with self.connect() as db:
            rows = db.execute(
                """
                SELECT *
                FROM scans
                WHERE resolution_state = 'pending'
                  AND next_resolve_epoch <= ?
                ORDER BY completed_at_utc
                LIMIT 50
                """,
                (now,),
            ).fetchall()
        for row in rows:
            deadline = parse_utc(
                row["resolve_deadline_utc"]
            )
            if deadline and deadline <= datetime.now(timezone.utc):
                with self.connect() as db:
                    db.execute(
                        """
                        UPDATE scans
                        SET resolution_state = 'unresolved',
                            updated_at_utc = ?
                        WHERE id = ?
                        """,
                        (
                            utc_now(),
                            row["id"],
                        ),
                    )
                continue
            result = smb.resolve(
                Path(row["original_path"]),
                row["relative_path"],
                row["file_mtime_utc"],
                row["completed_at_utc"],
            )
            client = norm_ip(
                result.get("client")
            )
            attempts = row["resolve_attempts"] + 1
            with self.connect() as db:
                if client:
                    db.execute(
                        """
                        UPDATE scans
                        SET scanner_ip = ?,
                            resolution_source = ?,
                            resolution_reason = ?,
                            resolution_confidence = ?,
                            resolution_record_id = ?,
                            resolution_state = 'resolved',
                            resolve_attempts = ?,
                            updated_at_utc = ?
                        WHERE id = ?
                        """,
                        (
                            client,
                            result["source"],
                            result["reason"],
                            result["confidence"],
                            result["record_id"],
                            attempts,
                            utc_now(),
                            row["id"],
                        ),
                    )
                else:
                    db.execute(
                        """
                        UPDATE scans
                        SET resolve_attempts = ?,
                            next_resolve_epoch = ?,
                            resolution_reason = ?,
                            updated_at_utc = ?
                        WHERE id = ?
                        """,
                        (
                            attempts,
                            time.time() + SMB_QUERY_INTERVAL_SECONDS,
                            result["reason"],
                            utc_now(),
                            row["id"],
                        ),
                    )

# -------------------------------------------------
# Aggregate completed scans by day and scanner IP address.
# -------------------------------------------------
    def data(self):
        with self.connect() as db:
            rows = db.execute(
                """
                SELECT
                    completed_date,
                    scanner_ip,
                    COUNT(*) AS files,
                    COALESCE(SUM(pages), 0) AS pages,
                    COALESCE(SUM(size_bytes), 0) AS bytes,
                    MIN(completed_at_local) AS first_scan,
                    MAX(completed_at_local) AS last_scan
                FROM scans
                WHERE resolution_state = 'resolved'
                  AND scanner_ip IS NOT NULL
                GROUP BY
                    completed_date,
                    scanner_ip
                ORDER BY
                    completed_date,
                    scanner_ip
                """
            ).fetchall()
        return [
            row
            for row in rows
            if norm_ip(row["scanner_ip"])
        ]
    @staticmethod
    def write_csv(path, fields, rows):

# -------------------------------------------------------------------------------------
# Write CSV data to a temporary file first, then atomically replace the report file.
# -------------------------------------------------------------------------------------

        path.parent.mkdir(
            parents=True,
            exist_ok=True,
        )
        tmp = path.with_suffix(
            path.suffix + ".tmp"
        )
        with tmp.open(
            "w",
            encoding="utf-8-sig",
            newline="",
        ) as csv_file:
            writer = csv.DictWriter(
                csv_file,
                fieldnames=fields,
                delimiter=";",
                extrasaction="ignore",
            )
            writer.writeheader()
            writer.writerows(rows)
            csv_file.flush()
            os.fsync(csv_file.fileno())
        os.replace(
            tmp,
            path,
        )
    def build_report_html(
        self,
        embedded,
        latest,
        interval_label,
    ):
        
# ---------------------------------
# Build the analytics HTML report.
# ---------------------------------
        generated = datetime.now().strftime(
            "%d.%m.%Y %H:%M:%S"
        )
        return f"""<!doctype html>
<html lang="en">
<head>
    <meta charset="utf-8">
    <meta name="viewport" content="width=device-width, initial-scale=1">
    <title>ScanProxy Analytics</title>
    <style>
        body {{
            font: 14px Segoe UI, Arial;
            background: #f3f5f8;
            color: #1f2937;
            margin: 0;
        }}

        main {{
            max-width: 1200px;
            margin: auto;
            padding: 24px;
        }}

        header,
        .ctrl,
        .cards {{
            display: flex;
            gap: 12px;
            flex-wrap: wrap;
            align-items: end;
        }}

        .card,
        table {{
            background: white;
            box-shadow: 0 2px 10px #0001;
        }}

        .card {{
            padding: 14px;
            border-radius: 12px;
            min-width: 140px;
        }}

        .v {{
            font-size: 23px;
            font-weight: 700;
        }}

        table {{
            width: 100%;
            border-collapse: collapse;
            margin-top: 16px;
        }}

        th,
        td {{
            padding: 11px;
            border-bottom: 1px solid #eee;
            text-align: left;
        }}

        .n {{
            text-align: right;
        }}

        select,
        input {{
            padding: 8px;
        }}

        small {{
            color: #667085;
        }}
    </style>
</head>
<body>
    <main>
        <header>
            <div>
                <h1>ScanProxy Analytics</h1>
                <small>
                    Unambiguous SMB client IP addresses only &middot;
                    Scanproxy {escape(VERSION)}
                </small>
            </div>
        </header>

        <div class="ctrl">
            <label>
                View
                <select id="m">
                    <option value="day">Day</option>
                    <option value="week">Week</option>
                    <option value="month">Month</option>
                    <option value="year">Year</option>
                </select>
            </label>

            <label>
                Date
                <input id="d" type="date" value="{latest}">
            </label>

            <b id="r"></b>
        </div>

        <div class="cards">
            <div class="card">
                Scanner
                <div class="v" id="c">0</div>
            </div>

            <div class="card">
                Files
                <div class="v" id="f">0</div>
            </div>

            <div class="card">
                Pages
                <div class="v" id="p">0</div>
            </div>

            <div class="card">
                Data
                <div class="v" id="b">0 MB</div>
            </div>

            <div class="card">
                First scan
                <div class="v" id="e">-</div>
            </div>

            <div class="card">
                Last scan
                <div class="v" id="l">-</div>
            </div>
        </div>

        <table>
            <thead>
                <tr>
                    <th>Scanner-IP</th>
                    <th class="n">Files</th>
                    <th class="n">Pages</th>
                    <th class="n">Data MB</th>
                    <th>Last scan</th>
                </tr>
            </thead>
            <tbody id="t"></tbody>
        </table>

        <small>
            Refreshed in batches every {interval_label}.
            Generated: {generated}
        </small>
    </main>

    <script>
        const D = {embedded};
        const m = document.getElementById('m');
        const d = document.getElementById('d');
        const t = document.getElementById('t');

        const pad = n => String(n).padStart(2, '0');

        const iso = x =>
            `${{x.getFullYear()}}-${{pad(x.getMonth() + 1)}}-${{pad(x.getDate())}}`;

        const pd = s => {{
            const a = s.split('-').map(Number);
            return new Date(a[0], a[1] - 1, a[2]);
        }};

        const fd = s => {{
            const x = pd(s);
            return `${{pad(x.getDate())}}.${{pad(x.getMonth() + 1)}}.${{x.getFullYear()}}`;
        }};

        const dt = s =>
            s
                ? new Date(s).toLocaleString(
                    'en-GB',
                    {{dateStyle: 'short', timeStyle: 'short'}}
                )
                : '-';

        function rr() {{
            const x = pd(d.value);
            let a = new Date(x);
            let z = new Date(x);

            if (m.value === 'week') {{
                a.setDate(x.getDate() - ((x.getDay() + 6) % 7));
                z = new Date(a);
                z.setDate(a.getDate() + 6);
            }}

            if (m.value === 'month') {{
                a = new Date(x.getFullYear(), x.getMonth(), 1);
                z = new Date(x.getFullYear(), x.getMonth() + 1, 0);
            }}

            if (m.value === 'year') {{
                a = new Date(x.getFullYear(), 0, 1);
                z = new Date(x.getFullYear(), 11, 31);
            }}

            const A = iso(a);
            const Z = iso(z);
            const q = D.filter(v => v.date >= A && v.date <= Z);
            const m0 = new Map();

            let F = 0;
            let P = 0;
            let B = 0;
            let L = '';
            let E = '';

            for (const v of q) {{
                F += v.files;
                P += v.pages;
                B += v.bytes;
                L = !L || v.last > L ? v.last : L;
                E = !E || v.first < E ? v.first : E;

                const o = m0.get(v.scanner) || {{
                    f: 0,
                    p: 0,
                    b: 0,
                    l: '',
                }};

                o.f += v.files;
                o.p += v.pages;
                o.b += v.bytes;
                o.l = !o.l || v.last > o.l ? v.last : o.l;

                m0.set(v.scanner, o);
            }}

            document.getElementById('r').textContent =
                `${{fd(A)}} to ${{fd(Z)}}`;

            document.getElementById('c').textContent = m0.size;
            document.getElementById('f').textContent = F;
            document.getElementById('p').textContent = P;

            document.getElementById('b').textContent =
                (B / 1048576).toLocaleString(
                    'en-GB',
                    {{minimumFractionDigits: 2, maximumFractionDigits: 2}}
                ) + ' MB';

            document.getElementById('e').textContent = dt(E);
            document.getElementById('l').textContent = dt(L);

            t.innerHTML = [...m0]
                .sort()
                .map(([s, o]) => `
                    <tr>
                        <td>${{s}}</td>
                        <td class="n">${{o.f}}</td>
                        <td class="n">${{o.p}}</td>
                        <td class="n">${{
                            (o.b / 1048576).toLocaleString(
                                'en-GB',
                                {{
                                    minimumFractionDigits: 2,
                                    maximumFractionDigits: 2,
                                }}
                            )
                        }}</td>
                        <td>${{dt(o.l)}}</td>
                    </tr>
                `)
                .join('') || '<tr><td colspan="5">No data</td></tr>';
        }}

        m.onchange = rr;
        d.onchange = rr;
        rr();
    </script>
</body>
</html>
"""

# ---------------------------------------------------------------------------------------------
# Build CSV and HTML analytics reports from SQLite and remove expired daily reports.
# ---------------------------------------------------------------------------------------------

    def refresh_reports(self):
        data = self.data()
        cutoff = (
            date.today()
            - timedelta(days=DAILY_RETENTION_DAYS)
        )
        by_day = {}
        groups = {
            "week": {},
            "month": {},
            "year": {},
        }
        for row in data:
            completed_date = date.fromisoformat(
                row["completed_date"]
            )
            by_day.setdefault(
                completed_date,
                [],
            ).append(row)
            item = {
                "scanner": row["scanner_ip"],
                "files": row["files"],
                "pages": row["pages"],
                "bytes": row["bytes"],
                "last": row["last_scan"],
            }
            week_start = (
                completed_date
                - timedelta(days=completed_date.weekday())
            )
            iso_week = week_start.isocalendar()

            month_start = date(
                completed_date.year,
                completed_date.month,
                1,
            )
            next_month = date(
                completed_date.year + (completed_date.month == 12),
                (completed_date.month % 12) + 1,
                1,
            )
            month_end = (
                next_month
                - timedelta(days=1)
            )
            keys = (
                (
                    "week",
                    f"{iso_week.year}-KW{iso_week.week:02d}-scanproxy.csv",
                    week_start,
                    week_start + timedelta(days=6),
                ),
                (
                    "month",
                    (
                        f"{completed_date.year:04d}-"
                        f"{completed_date.month:02d}-scanproxy.csv"
                    ),
                    month_start,
                    month_end,
                ),
                (
                    "year",
                    f"{completed_date.year:04d}-scanproxy.csv",
                    date(completed_date.year, 1, 1),
                    date(completed_date.year, 12, 31),
                ),
            )
            for report_type, name, start, end in keys:
                period = groups[report_type].setdefault(
                    (
                        name,
                        start,
                        end,
                    ),
                    {},
                )
                bucket = period.setdefault(
                    item["scanner"],
                    {
                        "files": 0,
                        "pages": 0,
                        "bytes": 0,
                        "last": "",
                    },
                )
                bucket["files"] += item["files"]
                bucket["pages"] += item["pages"]
                bucket["bytes"] += item["bytes"]
                bucket["last"] = max(
                    bucket["last"],
                    item["last"],
                )
        for completed_date, same_day in by_day.items():
            if completed_date < cutoff:
                continue
            rows = [
                {
                    "Date": completed_date.strftime("%d.%m.%Y"),
                    "Scanner": row["scanner_ip"],
                    "Files": row["files"],
                    "Pages": row["pages"],
                    "Data_MB": mb(row["bytes"]),
                    "Last_Scan": str(row["last_scan"])[11:16],
                }
                for row in same_day
            ]
            self.write_csv(
                DAILY
                / f"{completed_date.strftime('%d.%m.%Y')}-scanproxy.csv",
                (
                    "Date",
                    "Scanner",
                    "Files",
                    "Pages",
                    "Data_MB",
                    "Last_Scan",
                ),
                rows,
            )
        report_folders = (
            ("week", WEEKLY),
            ("month", MONTHLY),
            ("year", YEARLY),
        )
        for report_type, folder in report_folders:
            for (
                name,
                start,
                end,
            ), scanners in groups[report_type].items():
                rows = [
                    {
                        "From": start.strftime("%d.%m.%Y"),
                        "To": end.strftime("%d.%m.%Y"),
                        "Scanner": scanner,
                        "Files": values["files"],
                        "Pages": values["pages"],
                        "Data_MB": mb(values["bytes"]),
                        "Last_Scan": str(
                            values["last"]
                        )[0:16].replace("T", " "),
                    }
                    for scanner, values in sorted(
                        scanners.items()
                    )
                ]
                self.write_csv(
                    folder / name,
                    (
                        "From",
                        "To",
                        "Scanner",
                        "Files",
                        "Pages",
                        "Data_MB",
                        "Last_Scan",
                    ),
                    rows,
                )
        for report_file in DAILY.glob("*-scanproxy.csv"):
            try:
                report_date = datetime.strptime(
                    report_file.name[:10],
                    "%d.%m.%Y",
                ).date()
                if report_date < cutoff:
                    report_file.unlink()
            except ValueError:
                pass
        embedded_data = [
            {
                "date": row["completed_date"],
                "scanner": row["scanner_ip"],
                "files": row["files"],
                "pages": row["pages"],
                "bytes": row["bytes"],
                "first": row["first_scan"],
                "last": row["last_scan"],
            }
            for row in data
        ]
        embedded = json.dumps(
            embedded_data,
            ensure_ascii=False,
            separators=(",", ":"),
        )
        latest = (
            data[-1]["completed_date"]
            if data
            else date.today().isoformat()
        )
        if SMB_QUERY_INTERVAL_SECONDS % 60 == 0:
            interval_label = (
                f"{SMB_QUERY_INTERVAL_SECONDS // 60} minutes"
            )
        else:
            interval_label = (
                f"{SMB_QUERY_INTERVAL_SECONDS} seconds"
            )
        html = self.build_report_html(
            embedded,
            latest,
            interval_label,
        )
        tmp = REPORT.with_suffix(".tmp")
        tmp.write_text(
            html,
            encoding="utf-8",
        )
        os.replace(
            tmp,
            REPORT,
        )
        self.log.info(
            "ANALYTICS UPDATED | "
            "Interval=%ss | HTML=%s",
            SMB_QUERY_INTERVAL_SECONDS,
            REPORT,
        )

# ==============================================================================
# File routing
# ==============================================================================

# ----------------------------------------------------------------------------------------------------------------
# Determine the destination path and create a unique name for moves when a file name already exists.
# ----------------------------------------------------------------------------------------------------------------

def destination(target, relative, action):
    dst = target / relative
    dst.parent.mkdir(
        parents=True,
        exist_ok=True,
    )
    if not dst.exists() or action == "copy":
        return dst
    stamp = time.strftime("%Y%m%d_%H%M%S")
    number = 1
    while True:
        suffix = (
            ""
            if number == 1
            else f"_{number}"
        )
        candidate = dst.with_name(
            f"{dst.stem}_{stamp}{suffix}{dst.suffix}"
        )
        if not candidate.exists():
            return candidate
        number += 1

# -----------------------------------------------------------------------------------------------
# Check file age and SMB status, then copy the file to the destination.
# Data is written to a .part file first and finalized atomically.
# -----------------------------------------------------------------------------------------------

def transfer(
    src,
    dst,
    action,
    reason,
    label,
    smb,
    analytics,
    log,
):
    tmp = dst.with_name(
        dst.name
        + "."
        + uuid.uuid4().hex
        + ".part"
    )
    try:
        before = src.stat()
        if time.time() - before.st_mtime < MIN_AGE:
            return False
        if action == "copy" and dst.exists():
            return False
        if analytics.pending_move(src):
            return False
        try:
            relative = str(
                src.relative_to(ROOT)
            )
        except ValueError:
            relative = src.name
        mtime = datetime.fromtimestamp(
            before.st_mtime,
            tz=timezone.utc,
        ).isoformat().replace(
            "+00:00",
            "Z",
        )
        resolution = smb.resolve(
            src,
            relative,
            mtime,
            utc_now(),
        )
        if resolution.get("is_open"):
            return False
        current = src.stat()
        if (
            before.st_size,
            before.st_mtime_ns,
        ) != (
            current.st_size,
            current.st_mtime_ns,
        ):
            return False
        shutil.copy2(
            src,
            tmp,
        )
        after = src.stat()
        if (
            before.st_size,
            before.st_mtime_ns,
        ) != (
            after.st_size,
            after.st_mtime_ns,
        ):
            tmp.unlink(
                missing_ok=True
            )
            return False
        os.replace(
            tmp,
            dst,
        )
        committed = utc_now()
        if action == "move":
            try:
                src.unlink()
            except OSError as exc:
                try:
                    analytics.add_move(
                        src,
                        dst,
                        relative,
                        label,
                        before,
                        mtime,
                        resolution,
                        exc,
                    )
                except Exception:
                    try:
                        dst.unlink()
                    except OSError:
                        log.exception(
                            "CRITICAL | Destination could not be rolled back after failed move "
                            "| Destination=%s",
                            dst,
                        )
                return False
        analytics.add_scan(
            src,
            relative,
            dst,
            label,
            mtime,
            resolution,
            completed=committed,
        )
        log.info(
            "SUCCESS | File=%s | Destination=%s | "
            "Action=%s | Reason=%s",
            src,
            dst,
            action,
            reason,
        )
        return True
    except FileNotFoundError:
        return False
    except Exception:
        log.exception(
            "ERROR | File=%s | Destination=%s",
            src,
            dst,
        )
        return False
    finally:
        tmp.unlink(
            missing_ok=True
        )

# -------------------------------------------------------------------------
# Process all enabled Excel rules and pass matching files to transfer().
# -------------------------------------------------------------------------

def process_rules(
    rules,
    smb,
    analytics,
    log,
):
    """
    What:
        Iterate through all enabled Excel rules, find matching files, and pass
        them to transfer(). The id and subfolder modes are supported.

    Why:
        File discovery and safe file transfer stay separate so rule definitions
        can change without duplicating transfer logic.
    """
    for (
        row,
        code,
        source,
        target,
        mode,
        enabled,
        action,
        scanner,
    ) in rules:
        if STOP.is_set():
            return
        if not active(enabled):
            continue
        if not source or not target:
            log.error(
                "RULE ERROR | Row=%s | Source/target is missing",
                row,
            )
            continue
        src = p(source)
        dst = p(target)
        mode = str(
            mode or ""
        ).strip().lower()
        action = str(
            action or "move"
        ).strip().lower()

        if action not in {
            "move",
            "copy",
        } or not src.is_dir():
            log.error(
                "RULE ERROR | Row=%s | "
                "Source/action is invalid",
                row,
            )
            continue
        if mode == "id":
            key = str(
                code or ""
            ).strip().lower()
            files = (
                file
                for file in src.iterdir()
                if file.is_file()
                and not file.name.endswith(".part")
                and (
                    file.stem.lower() == key
                    or file.name.lower().startswith(
                        key + "-"
                    )
                )
            )
        elif mode in {
            "subfolder",
            "subfolders",
        }:
            if dst == src or dst.is_relative_to(src):
                log.error(
                    "RULE ERROR | Target is inside source | Row=%s",
                    row,
                )
                continue
            files = (
                file
                for file in src.rglob("*")
                if file.is_file()
                and not file.name.endswith(".part")
            )
        else:
            log.error(
                "RULE ERROR | Row=%s | Mode=%s",
                row,
                mode,
            )
            continue
        for file in files:
            if mode == "id":
                relative = Path(file.name)
            else:
                relative = file.relative_to(src)
            label = str(
                scanner
                or code
                or (
                    relative.parts[0]
                    if len(relative.parts) > 1
                    else ""
                )
            )
            transfer(
                file,
                destination(
                    dst,
                    relative,
                    action,
                ),
                action,
                f"Excel row {row}",
                label,
                smb,
                analytics,
                log,
            )


# ==============================================================================
# Process control
# ==============================================================================


# -----------------------------------------------------------------------------------------
# Create a global Windows mutex and detect whether another instance is already running.
# -----------------------------------------------------------------------------------------

def single_instance(log):
    if os.name != "nt":
        return True
    kernel32 = ctypes.WinDLL(
        "kernel32",
        use_last_error=True,
    )
    kernel32.CreateMutexW.restype = ctypes.c_void_p
    kernel32.CreateMutexW.argtypes = (
        ctypes.c_void_p,
        ctypes.c_bool,
        ctypes.c_wchar_p,
    )
    kernel32.CloseHandle.argtypes = (
        ctypes.c_void_p,
    )
    ctypes.set_last_error(0)
    handle = kernel32.CreateMutexW(
        None,
        False,
        "Global\\ScanProxy",
    )
    already_running = (
        ctypes.get_last_error() == 183
    )
    if not handle or already_running:
        if handle:
            kernel32.CloseHandle(handle)
        log.error(
            "ABORT | ScanProxy is already running"
        )
        return False
    return handle

# ----------------------------------------------------------------------------------
# Initialize logging, single-instance protection, analytics, and SMB handling.
# The main loop uses separate intervals for file scanning and analytics.
# ----------------------------------------------------------------------------------

def main():
    log, analytics_log = get_logs()
    mutex = single_instance(log)

    if mutex is False:
        return
    signal.signal(
        signal.SIGINT,
        lambda *_: STOP.set(),
    )
    analytics = Analytics(
        analytics_log
    )
    smb = Smb(log)
    analytics.init()
    last_config_mtime = None
    rules = []
    sheet = ""
    next_file_scan = 0.0
    next_analytics = (
        time.monotonic()
        + SMB_QUERY_INTERVAL_SECONDS
    )
    log.info(
        "START | Version=%s | Root=%s | File scan=%ss | "
        "SMB query=%ss | SMB lookback=%smin | Ignored=%s",
        VERSION,
        ROOT,
        FILE_SCAN_INTERVAL_SECONDS,
        SMB_QUERY_INTERVAL_SECONDS,
        SMB_EVENT_LOOKBACK_MINUTES,
        ",".join(SMB_IGNORED_NETWORKS),
    )
    analytics_log.info(
        "START | Version=%s | Analytics=%s | SMB query=%ss",
        VERSION,
        ANALYTICS,
        SMB_QUERY_INTERVAL_SECONDS,
    )
    print(
        f"ScanProxy {VERSION} is running. Press Ctrl+C to stop.",
        flush=True,
    )
    try:
        while not STOP.is_set():
            try:
                analytics.retry_moves()
                now = time.monotonic()
                config_mtime = CONFIG.stat().st_mtime_ns
                # Reload Excel only when the file actually changes.
                # After a change, allow the next file scan immediately so new rules do not wait for the interval.
                if config_mtime != last_config_mtime:
                    sheet, rules = load_rules()
                    last_config_mtime = config_mtime
                    next_file_scan = 0.0
                    log.info(
                        "EXCEL LOADED | Sheet=%s | Rules=%s",
                        sheet,
                        len(rules),
                    )
                # File discovery is the most CPU-intensive step, so it runs only at FILE_SCAN_INTERVAL.
                if now >= next_file_scan:
                    process_rules(
                        rules,
                        smb,
                        analytics,
                        log,
                    )
                    next_file_scan = (
                        time.monotonic()
                        + FILE_SCAN_INTERVAL_SECONDS
                    )
                # SMB re-resolution and report rebuilding are batched so PowerShell, SQLite, and HTML generation do not run continuously.
                if now >= next_analytics:
                    analytics.resolve_pending(smb)
                    analytics.refresh_reports()
                    next_analytics = (
                        time.monotonic()
                        + SMB_QUERY_INTERVAL_SECONDS
                    )
            except Exception:
                log.exception(
                    "RUN ERROR"
                )
            STOP.wait(
                MAIN_LOOP_SLEEP_SECONDS
            )

    finally:
        # Try one final analytics run during shutdown.
        try:
            analytics.resolve_pending(smb)
            analytics.refresh_reports()
        except Exception:
            analytics_log.exception(
                "ANALYTICS ERROR | Shutdown"
            )
        if os.name == "nt" and mutex not in (
            True,
            False,
        ):
            ctypes.WinDLL(
                "kernel32"
            ).CloseHandle(
                ctypes.c_void_p(mutex)
            )
        analytics_log.info("END")
        log.info("END")
        print(
            "ScanProxy stopped.",
            flush=True,
        )
if __name__ == "__main__":
    main()